"""The campaign's strategic briefing, as a PPTX (FEAT-CAMPAIGN-BRIEFING-DECK ticket 04).

A military situation brief after the VEAF briefing template — 16:9, white, bold title top left, text
left-aligned and never justified, Calibri — whose facts are generated from the campaign (the map,
the intelligence, the objectives, the annex of rules) around the prose of `briefing.yaml`.

Only the annex speaks of the game's mechanics; everything else is said as a staff would say it.
A page too full for its frame is continued on the next one, never shrunk: the text is wrapped with
the font's real metrics, as PowerPoint wraps it, to know where a page ends.
"""

from __future__ import annotations

import functools
from dataclasses import dataclass, field
from pathlib import Path

from PIL import Image, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.presentation import Presentation as PresentationType
from pptx.slide import Slide
from pptx.text.text import TextFrame, _Paragraph
from pptx.util import Emu, Pt
from veaf_libs.i18n import t
from veaf_libs.map_tiles import Fetch

from campaign_manager.briefing_prose import BriefingProse
from campaign_manager.intelligence import enemy_picture, friendly_picture, reserve_text
from campaign_manager.models import COALITIONS, CampaignDefinition, CampaignState, CampaignZone
from campaign_manager.strategic_map import OWNER_COLOURS, MapReport, render_strategic_map
from campaign_manager.turn_manager import enemy_of

#: The deck, in the mission's sub-folder (`missions/mission-NN/`).
DECK_FILE = "briefing-campagne.pptx"

#: The map the deck shows, next to it.
MAP_FILE = "carte-strategique.png"

#: The template's font, one Google Slides knows.
FONT = "Calibri"

_WIDTH, _HEIGHT = Emu(12192000), Emu(6858000)
_MARGIN = Emu(457200)
_BODY_TOP = Emu(1150000)
_BODY_BOTTOM = Emu(380000)
_BODY_SIZE = 14
_HEADING_SIZE = 16
_SPACE_BEFORE_HEADING = 8
_LINE_SPACING = 1.2
_BLUE = RGBColor(0x1F, 0x5F, 0xBF)
_GREY = RGBColor(0x55, 0x55, 0x55)


@dataclass(frozen=True)
class Block:
    """A heading and its paragraphs; either may be empty."""

    heading: str | None
    paragraphs: tuple[str, ...]
    colour: tuple[int, int, int] | None = None


@dataclass
class Page:
    """A page of blocks, under its title."""

    title: str
    blocks: list[Block] = field(default_factory=list)
    note: str | None = None
    """A line at the foot of the page, small and grey."""
    is_map: bool = False
    """The strategic map's page, drawn with its picture and legend."""


@dataclass(frozen=True)
class DeckReport:
    """What generating the deck produced."""

    path: Path
    map: MapReport
    pages: int
    prose: bool
    """Whether the written half was there; without it the deck is the generated half only."""


# ---------------------------------------------------------------------------
# Measuring text: where a page ends
# ---------------------------------------------------------------------------


@functools.lru_cache(maxsize=8)
def _measuring_font(size: int, bold: bool) -> ImageFont.FreeTypeFont | None:
    """The template's font at a size, or a wider free font in its place; ``None`` without either."""
    for name in ("calibrib.ttf", "DejaVuSans-Bold.ttf") if bold else ("calibri.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size * 10)
        except OSError:
            continue
    return None


def _text_width(text: str, size: int, bold: bool) -> float:
    """The width of a text in points, at a font size in points."""
    font = _measuring_font(size, bold)
    if font is None:
        return len(text) * size * 0.55  # wider than Calibri: errs towards a new page
    return font.getlength(text) / 10


def line_count(text: str, size: int, width: float, bold: bool = False) -> int:
    """How many lines a paragraph takes in a frame that wide, wrapped at the spaces.

    Args:
        text: The paragraph.
        size: Its font size, in points.
        width: The frame's width, in points.
        bold: Whether it is bold.

    Returns:
        The number of lines, 1 at least.
    """
    lines, current = 1, ""
    for word in text.split():
        candidate = f"{current} {word}" if current else word
        if current and _text_width(candidate, size, bold) > width:
            lines, current = lines + 1, word
        else:
            current = candidate
    return lines


def block_height(block: Block, width: float) -> float:
    """The height of a block, in points, in a frame that wide."""
    height = 0.0
    if block.heading:
        height += (
            _SPACE_BEFORE_HEADING
            + line_count(block.heading, _HEADING_SIZE, width, bold=True) * _HEADING_SIZE * _LINE_SPACING
        )
    for paragraph in block.paragraphs:
        height += line_count(paragraph, _BODY_SIZE, width) * _BODY_SIZE * _LINE_SPACING
    return height


def _body_width() -> float:
    return Emu(_WIDTH - 2 * _MARGIN).pt - 14  # the text frame's own inner margins


def _body_height(note: bool) -> float:
    return Emu(_HEIGHT - _BODY_TOP - _BODY_BOTTOM).pt - (28 if note else 0)


def paginate(page: Page) -> list[Page]:
    """Split a page whose blocks do not fit into pages that do, continued under the same title.

    A block is kept whole when it fits on a page of its own, and split between its paragraphs
    otherwise; a heading never ends a page.

    Args:
        page: The page as written.

    Returns:
        One page or more.
    """
    width, height = _body_width(), _body_height(page.note is not None)
    pages = [Page(page.title, note=page.note)]
    used = 0.0
    for block in page.blocks:
        pieces = [block]
        if block_height(block, width) > height:
            pieces = [Block(block.heading, block.paragraphs[:1], block.colour)] + [
                Block(None, (paragraph,), block.colour) for paragraph in block.paragraphs[1:]
            ]
        for piece in pieces:
            needed = block_height(piece, width)
            if used + needed > height and pages[-1].blocks:
                pages.append(Page(t("campaign.deck.continued", title=page.title), note=page.note))
                used = 0.0
            pages[-1].blocks.append(piece)
            used += needed
    return pages


# ---------------------------------------------------------------------------
# What the pages say
# ---------------------------------------------------------------------------


def _join(items: list[str]) -> str:
    """Items said as a list in words: `a`, `a et b`, `a, b et c`."""
    return items[0] if len(items) == 1 else t("campaign.deck.and", first=", ".join(items[:-1]), last=items[-1])


def _labels(campaign: CampaignDefinition, names: tuple[str, ...] | list[str]) -> str:
    return _join([campaign.zone(name).label for name in names])


def _zones_of(campaign: CampaignDefinition, state: CampaignState, side: str) -> list[CampaignZone]:
    return [zone for zone in campaign.zones if state.zones[zone.name].owner == side]


def _prose_block(prose: BriefingProse | None, heading_key: str, section: str, part: str) -> list[Block]:
    paragraphs = prose.text(section, part) if prose else ()
    return [Block(t(heading_key), paragraphs)] if paragraphs else []


def _strategic_page(prose: BriefingProse | None) -> Page | None:
    blocks = _prose_block(prose, "campaign.deck.heading.political", "situation", "political")
    blocks += _prose_block(prose, "campaign.deck.heading.economic", "situation", "economic")
    return Page(t("campaign.deck.title.strategic"), blocks) if blocks else None


def _military_page(campaign: CampaignDefinition, state: CampaignState, prose: BriefingProse | None) -> Page:
    player = campaign.player_side
    enemy = enemy_of(player)
    enemy_lines = [
        f"– {zone.label} : {enemy_picture(campaign, state, zone)}" for zone in _zones_of(campaign, state, enemy)
    ]
    depots = [zone.name for zone in _zones_of(campaign, state, enemy) if zone.kind == "logistics"]
    enemy_lines.append(
        t("campaign.deck.enemy_reserve_fed", depots=_labels(campaign, depots))
        if depots
        else t("campaign.deck.enemy_reserve")
    )
    for paragraph in prose.text("situation", "enemy_course_of_action") if prose else ():
        enemy_lines.append(t("campaign.deck.enemy_course_of_action", text=paragraph))
    friendly_lines = [
        f"– {zone.label} : {friendly_picture(campaign, zone)}." for zone in _zones_of(campaign, state, player)
    ]
    friendly_lines.append(t("campaign.deck.friendly_reserve", reserve=reserve_text(state.sides[player].reserve)))
    friendly_lines += [f"– {paragraph}" for paragraph in (prose.text("situation", "friendly") if prose else ())]
    blocks = [
        Block(t("campaign.deck.heading.enemy"), tuple(enemy_lines), OWNER_COLOURS[enemy]),
        Block(t("campaign.deck.heading.friendly"), tuple(friendly_lines), OWNER_COLOURS[player]),
    ]
    neutral = [t("campaign.deck.neutral_zone", zone=zone.label) for zone in _zones_of(campaign, state, "neutral")]
    if neutral:
        blocks.append(Block(t("campaign.deck.heading.neutral"), tuple(neutral)))
    return Page(t("campaign.deck.title.military"), blocks, note=t("campaign.deck.intel_note"))


def _intent_page(prose: BriefingProse | None) -> Page | None:
    blocks = _prose_block(prose, "campaign.deck.heading.mission", "mission", "mission")
    intent = [
        t(f"campaign.deck.intent.{part}", text=paragraph)
        for part in ("purpose", "main_effect", "method", "end_state")
        for paragraph in (prose.text("intent", part) if prose else ())
    ]
    if intent:
        blocks.append(Block(t("campaign.deck.heading.intent"), tuple(intent)))
    return Page(t("campaign.deck.title.intent"), blocks) if blocks else None


def _objectives_page(campaign: CampaignDefinition, state: CampaignState, prose: BriefingProse | None) -> Page:
    blocks = []
    for part in ("political", "military", "economic"):
        paragraphs = prose.text("objectives", part) if prose else ()
        if paragraphs:
            blocks.append(Block(t(f"campaign.deck.heading.objectives_{part}"), tuple(f"– {p}" for p in paragraphs)))
    conditions = []
    for objective in campaign.objectives:
        key = "capture" if objective.kind == "capture" else "destroy"
        conditions.append(t(f"campaign.deck.victory.{key}", zones=_labels(campaign, objective.zones)))
    left = max(campaign.missions - state.mission, 0)
    victory = (
        t("campaign.deck.victory.by", missions=campaign.missions),
        *conditions,
        t("campaign.deck.victory.left", left=left),
    )
    blocks.append(Block(t("campaign.deck.heading.victory"), victory))
    return Page(t("campaign.deck.title.objectives"), blocks)


def _concept_page(prose: BriefingProse | None) -> Page | None:
    if prose is None:
        return None
    blocks = [Block(phase.title, phase.text) for phase in prose.phases]
    attention = prose.text("concept", "attention")
    if attention:
        blocks.append(Block(t("campaign.deck.heading.attention"), tuple(f"– {p}" for p in attention)))
    return Page(t("campaign.deck.title.concept"), blocks) if blocks else None


def _roe_page(prose: BriefingProse | None) -> Page | None:
    blocks = []
    for part in ("targeting", "civilians", "self_defence"):
        paragraphs = prose.text("rules_of_engagement", part) if prose else ()
        if paragraphs:
            blocks.append(Block(t(f"campaign.deck.heading.roe_{part}"), tuple(f"– {p}" for p in paragraphs)))
    return Page(t("campaign.deck.title.roe"), blocks) if blocks else None


def _mission_page(prose: BriefingProse | None, mission: int) -> Page | None:
    page = prose.missions.get(mission) if prose else None
    if page is None:
        return None
    blocks = [Block(f"{index}. {task.title}", task.text) for index, task in enumerate(page.tasks, 1)]
    return Page(t("campaign.deck.title.mission", mission=mission, title=page.title), blocks)


def _annex_page(campaign: CampaignDefinition, state: CampaignState) -> Page:
    rules = campaign.rules
    output = rules.logistics_output
    capture = (
        t("campaign.deck.annex.neutral"),
        t("campaign.deck.annex.capture", minutes=max(campaign.capture_seconds // 60, 1)),
        t("campaign.deck.annex.garrison"),
    )
    between = [
        t(
            "campaign.deck.annex.logistics",
            armor=output.get("armor", 0),
            air_defense=output.get("air_defense", 0),
            transport=output.get("transport", 0),
        ),
        t("campaign.deck.annex.repairs", repairs=rules.repairs_per_mission),
    ]
    if rules.counter_attack:
        between.append(t("campaign.deck.annex.counter_attack"))
        for zone in _zones_of(campaign, state, "neutral"):
            bordering = {state.zones[name].owner for name in campaign.neighbours(zone.name)} & set(COALITIONS)
            if len(bordering) > 1:
                between.append(t("campaign.deck.annex.contested", zone=zone.label))
    else:
        between.append(t("campaign.deck.annex.no_counter_attack"))
    return Page(
        t("campaign.deck.title.annex"),
        [
            Block(t("campaign.deck.heading.annex_capture"), capture),
            Block(t("campaign.deck.heading.annex_between"), tuple(between)),
            Block(t("campaign.deck.heading.annex_continuity"), (t("campaign.deck.annex.continuity"),)),
        ],
    )


# ---------------------------------------------------------------------------
# Writing the deck
# ---------------------------------------------------------------------------


def _frame(slide: Slide, left: int, top: int, width: int, height: int) -> TextFrame:
    frame = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height)).text_frame
    frame.word_wrap = True
    return frame


def _run(
    paragraph: _Paragraph,
    text: str,
    size: int,
    *,
    bold: bool = False,
    italic: bool = False,
    colour: RGBColor | None = None,
) -> None:
    run = paragraph.add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold, run.font.italic = FONT, Pt(size), bold, italic
    if colour is not None:
        run.font.color.rgb = colour


def _titled_slide(deck: PresentationType, title: str) -> Slide:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    _run(_frame(slide, _MARGIN, Emu(300000), _WIDTH - 2 * _MARGIN, Emu(700000)).paragraphs[0], title, 30, bold=True)
    return slide


def _write_page(deck: PresentationType, page: Page) -> None:
    slide = _titled_slide(deck, page.title)
    frame = _frame(slide, _MARGIN, _BODY_TOP, _WIDTH - 2 * _MARGIN, _HEIGHT - _BODY_TOP - _BODY_BOTTOM)
    first = True
    for block in page.blocks:
        if block.heading:
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            paragraph.space_before = Pt(_SPACE_BEFORE_HEADING)
            _run(
                paragraph,
                block.heading,
                _HEADING_SIZE,
                bold=True,
                colour=RGBColor(*block.colour) if block.colour else None,
            )
        for text in block.paragraphs:
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            _run(paragraph, text, _BODY_SIZE)
    if page.note:
        note = _frame(slide, _MARGIN, _HEIGHT - Emu(600000), _WIDTH - 2 * _MARGIN, Emu(400000))
        _run(note.paragraphs[0], page.note, 11, italic=True, colour=_GREY)


def _write_cover(
    deck: PresentationType, campaign: CampaignDefinition, prose: BriefingProse | None, mission: int, contents: list[str]
) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = _frame(slide, _MARGIN, Emu(1700000), Emu(6600000), Emu(3000000))
    _run(frame.paragraphs[0], t("campaign.deck.operation"), 40, bold=True)
    _run(
        frame.add_paragraph(),
        (prose.operation if prose and prose.operation else campaign.name).upper(),
        60,
        bold=True,
        colour=_BLUE,
    )
    _run(
        frame.add_paragraph(),
        prose.subtitle if prose and prose.subtitle else t("campaign.deck.subtitle"),
        22,
        colour=_GREY,
    )
    _run(
        frame.add_paragraph(),
        t("campaign.deck.cover_line", theatre=campaign.theatre, missions=campaign.missions, mission=mission),
        16,
        italic=True,
        colour=_GREY,
    )
    if prose is None:
        _run(frame.add_paragraph(), t("campaign.deck.prose_missing"), 14, italic=True, colour=_GREY)
    toc = _frame(slide, Emu(7700000), Emu(1500000), Emu(4100000), Emu(4600000))
    _run(toc.paragraphs[0], t("campaign.deck.contents"), 20, bold=True)
    for item in contents:
        _run(toc.add_paragraph(), item, 15)


def _write_map(deck: PresentationType, title: str, report: MapReport) -> None:
    slide = _titled_slide(deck, title)
    with Image.open(report.path) as image:
        ratio = image.width / image.height
    height = _HEIGHT - Emu(1250000)
    width = min(int(height * ratio), Emu(8200000))
    height = int(width / ratio)
    slide.shapes.add_picture(str(report.path), _MARGIN, Emu(1050000), width=Emu(width), height=Emu(height))
    legend = _frame(slide, _MARGIN + width + Emu(300000), Emu(1200000), _WIDTH - width - 3 * _MARGIN, Emu(4500000))
    _run(legend.paragraphs[0], t("campaign.deck.legend"), 18, bold=True)
    for owner in ("blue", "red", "neutral"):
        paragraph = legend.add_paragraph()
        _run(paragraph, "● ", 18, colour=RGBColor(*OWNER_COLOURS[owner]))
        _run(paragraph, t(f"campaign.deck.legend_{owner}"), 15)
    _run(legend.add_paragraph(), t("campaign.deck.legend_axes"), 15)
    if report.offline:
        _run(legend.add_paragraph(), t("campaign.deck.map_offline"), 12, italic=True, colour=_GREY)


def campaign_deck(
    campaign: CampaignDefinition,
    state: CampaignState,
    prose: BriefingProse | None,
    folder: Path,
    *,
    cache_dir: Path | None = None,
    fetch: Fetch | None = None,
) -> DeckReport:
    """Write the strategic briefing of the coming mission, and its map, into a folder.

    Args:
        campaign: The validated campaign.
        state: The campaign state the coming mission starts from.
        prose: The written half (`briefing.yaml`), or ``None`` for the generated half only.
        folder: The mission's sub-folder, `missions/mission-NN/`.
        cache_dir: Where map tiles are kept between runs.
        fetch: How a map tile is downloaded (tests pass their own).

    Returns:
        Where the deck is, and what it was made from.
    """
    mission = state.mission + 1
    folder.mkdir(parents=True, exist_ok=True)
    map_report = render_strategic_map(campaign, state, folder / MAP_FILE, cache_dir=cache_dir, fetch=fetch)

    numbered: list[Page | None] = [
        _strategic_page(prose),
        _military_page(campaign, state, prose),
        Page(t("campaign.deck.title.map"), is_map=True),
        _intent_page(prose),
        _objectives_page(campaign, state, prose),
        _concept_page(prose),
        _roe_page(prose),
        _mission_page(prose, mission),
    ]
    sections = [page for page in numbered if page is not None]
    for number, page in enumerate(sections, 1):
        page.title = f"{number}. {page.title}"
    annex = _annex_page(campaign, state)
    contents = [page.title for page in sections] + [annex.title]

    deck = Presentation()
    deck.slide_width, deck.slide_height = _WIDTH, _HEIGHT
    _write_cover(deck, campaign, prose, mission, contents)
    for page in [*sections, annex]:
        if page.is_map:
            _write_map(deck, page.title, map_report)
            continue
        for part in paginate(page):
            _write_page(deck, part)
    out = folder / DECK_FILE
    deck.save(str(out))
    return DeckReport(path=out, map=map_report, pages=len(deck.slides), prose=prose is not None)
