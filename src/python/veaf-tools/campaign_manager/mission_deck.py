"""The coming campaign mission's own briefing, as a PPTX (FEAT-CAMPAIGN-MISSION-BRIEFING ticket 04).

The VEAF mission briefing (`.prompts/new-objective-mission.fr.md` §3 and §6), page for page as far as
a campaign mission allows, next to the campaign's strategic one and written with its helpers — the
VEAF template, pagination by the font's metrics, the military register.

Everything comes from the campaign, `briefing.yaml` (the mission's title and tasks) and the **built**
mission (date, time, weather, bullseye, flights, support, carrier, QRA zones). What a campaign
mission cannot have: target coordinates and a flight plan — a garrison is drawn when the mission
starts, so the briefing gives the zones' centres and says positions are found in flight.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.util import Emu
from veaf_libs.coordinates import latlon_to_xy, to_dms, xy_to_latlon
from veaf_libs.i18n import t
from veaf_libs.map_tiles import Fetch
from veaf_libs.terrain_elevation import grid_for_theatre, metres_to_feet

from campaign_manager.briefing_deck import (
    _BLUE,
    _GREY,
    _HEIGHT,
    _MARGIN,
    _WIDTH,
    Block,
    Page,
    _after_colon,
    _frame,
    _join,
    _run,
    _titled_slide,
    _write_page,
    _zone_line,
    paginate,
)
from campaign_manager.briefing_prose import BriefingProse, MissionPage
from campaign_manager.intelligence import enemy_picture
from campaign_manager.mission_picture import NM, Carrier, MissionPicture
from campaign_manager.models import CampaignDefinition, CampaignState, CampaignZone
from campaign_manager.strategic_map import zone_position
from campaign_manager.tactical_map import TACTICAL_MAP_FILE, render_objective_map, render_tactical_map, zoom_file
from campaign_manager.turn_manager import enemy_of

#: The deck, in the mission's sub-folder (`missions/mission-NN/`).
MISSION_DECK_FILE = "briefing-mission.pptx"

#: The guard frequency every flight monitors, MHz.
_GUARD = 243.0

#: hPa per mmHg.
_HPA = 1.33322


@dataclass(frozen=True)
class MissionDeckReport:
    """What generating the mission briefing produced."""

    path: Path
    pages: int
    maps: list[Path]
    offline: bool
    """Whether a map had to do without its background."""


def _names(title: str, zone: CampaignZone) -> bool:
    """Whether a task's title names a zone: its name or the name the players read, whole words, case aside."""
    return any(re.search(rf"\b{re.escape(name)}\b", title, re.IGNORECASE) for name in {zone.name, zone.label})


def objective_zones(campaign: CampaignDefinition, page: MissionPage | None) -> list[CampaignZone]:
    """The zones the coming mission's tasks name, in the tasks' order, or the campaign's objectives.

    A task names a zone when its title holds the zone's name or the name the players read, whole
    words, case aside: "Frapper le dépôt de Khobi" names "Dépôt de Khobi".

    Args:
        campaign: The validated campaign.
        page: The coming mission's page of `briefing.yaml`, or ``None``.

    Returns:
        The zones, each once.
    """
    found: list[CampaignZone] = []
    for task in page.tasks if page else ():
        for zone in campaign.zones:
            if zone not in found and _names(task.title, zone):
                found.append(zone)
    if not found:
        for objective in campaign.objectives:
            found += [campaign.zone(name) for name in objective.zones if campaign.zone(name) not in found]
    return found


def _line(name: str, text: str) -> str:
    """``name: text``, punctuated as the language wants it."""
    return t("campaign.mission_deck.line", name=name, text=text)


def _clock(seconds: int) -> str:
    return f"{seconds // 3600:02d}h{seconds % 3600 // 60:02d}"


def _mhz(value: float | None) -> str:
    return "?" if value is None else f"{value:.3f}".rstrip("0").rstrip(".") if value % 1 else f"{value:.1f}"


def _zone_xy(campaign: CampaignDefinition, zone: CampaignZone) -> tuple[float, float]:
    return latlon_to_xy(campaign.theatre, *zone_position(campaign, zone))


def _nearest_zone(campaign: CampaignDefinition, x: float, y: float) -> CampaignZone:
    return min(campaign.zones, key=lambda zone: math.dist(_zone_xy(campaign, zone), (x, y)))


def _reference(campaign: CampaignDefinition, picture: MissionPicture) -> tuple[str, tuple[float, float]] | None:
    """The friendly point the bullseye is given from: the first airfield with slots, else the carrier."""
    for airfield in picture.airfields:
        for zone in campaign.zones:
            if zone.location.airfield == airfield.name:
                return airfield.name, _zone_xy(campaign, zone)
    if picture.carriers:
        return picture.carriers[0].name, (picture.carriers[0].x, picture.carriers[0].y)
    return None


def _status(campaign: CampaignDefinition, state: CampaignState, zone: CampaignZone) -> str:
    owner = state.zones[zone.name].owner
    if owner == campaign.player_side:
        return t("campaign.mission_deck.zone_ours")
    if owner == "neutral":
        return t("campaign.mission_deck.zone_neutral")
    return enemy_picture(campaign, state, zone)


def _weather_lines(picture: MissionPicture) -> tuple[str, ...]:
    weather = picture.weather
    sky = (
        t("campaign.mission_deck.sky_clear")
        if weather.cover == "clear"
        else t(
            "campaign.mission_deck.sky_clouds",
            cover=t(f"campaign.mission_deck.cover.{weather.cover}"),
            base=round(metres_to_feet(weather.cloud_base), -2),
        )
    )
    visibility = (
        t("campaign.mission_deck.visibility_10")
        if weather.visibility >= 9000
        else t("campaign.mission_deck.visibility_km", km=round(weather.visibility / 1000))
    )
    extras = [t("campaign.mission_deck.rain")] if weather.rain else []
    extras += [t("campaign.mission_deck.fog")] if weather.fog else []
    winds = weather.winds
    return (
        t(
            "campaign.mission_deck.weather",
            sky=sky,
            visibility=visibility,
            temperature=round(weather.temperature),
            qnh=round(weather.qnh),
            hpa=round(weather.qnh * _HPA),
        ),
        t(
            "campaign.mission_deck.wind",
            ground=f"{winds['ground'].origin:03d}° / {winds['ground'].knots} kt",
            at2000=f"{winds['2000'].origin:03d}° / {winds['2000'].knots} kt",
            at8000=f"{winds['8000'].origin:03d}° / {winds['8000'].knots} kt",
        ),
        *extras,
    )


def _qra_line(campaign: CampaignDefinition, picture: MissionPicture) -> str | None:
    if not picture.qra_zones:
        return None
    near = _join([_nearest_zone(campaign, zone.x, zone.y).label for zone in picture.qra_zones])
    line = t("campaign.mission_deck.qra", zones=near)
    # said as intelligence: the tiers and the level are the mission maker's, never the players'
    if any(zone.scaled for zone in picture.qra_zones):
        line += " " + t("campaign.mission_deck.qra_scaled")
    return line


def _situation_page(
    campaign: CampaignDefinition,
    state: CampaignState,
    prose: BriefingProse | None,
    page: MissionPage | None,
    picture: MissionPicture,
    mission: int,
) -> Page:
    context = list(prose.text("mission", "mission")) if prose else []
    if prose and len(prose.phases) >= mission:
        context += list(prose.phases[mission - 1].text)
    blocks = [Block(t("campaign.mission_deck.heading.context"), tuple(context))] if context else []
    if page and page.tasks:
        tasks = tuple(f"– {task.title.split(' — ')[0]}" for task in page.tasks)
        blocks.append(Block(t("campaign.mission_deck.heading.mission"), tasks))
    lat, lon = xy_to_latlon(campaign.theatre, *picture.bullseye)
    bullseye = to_dms(lat, lon)
    reference = _reference(campaign, picture)
    if reference:
        name, (x, y) = reference
        dx, dy = picture.bullseye[0] - x, picture.bullseye[1] - y
        bearing = round(math.degrees(math.atan2(dy, dx))) % 360 or 360
        bullseye = t(
            "campaign.mission_deck.bullseye",
            dms=bullseye,
            bearing=f"{bearing:03d}",
            range=round(math.hypot(dx, dy) / NM),
            base=name,
        )
    blocks.append(Block(t("campaign.mission_deck.heading.bullseye"), (bullseye,)))
    departures = [
        t("campaign.mission_deck.departure_carrier", carrier=c.name, heading=f"{c.heading:03d}")
        for c in picture.carriers
    ]
    departures += [t("campaign.mission_deck.departure_airfield", airfield=a.name) for a in picture.airfields]
    if departures:
        blocks.append(Block(t("campaign.mission_deck.heading.departure"), tuple(departures)))
    enemy = enemy_of(campaign.player_side)
    threats = [
        _zone_line(zone, enemy_picture(campaign, state, zone))
        for zone in campaign.zones
        if state.zones[zone.name].owner == enemy
    ]
    qra = _qra_line(campaign, picture)
    if qra:
        threats.insert(0, f"– {qra}")
    if threats:
        blocks.append(Block(t("campaign.mission_deck.heading.threat"), tuple(threats)))
    blocks.append(Block(t("campaign.mission_deck.heading.weather"), _weather_lines(picture)))
    date = t("campaign.mission_deck.date", day=picture.date.day, month=picture.date.month, year=picture.date.year)
    blocks.append(
        Block(
            t("campaign.mission_deck.heading.time"),
            (t("campaign.mission_deck.time", date=date, time=_clock(picture.start_time)),),
        )
    )
    return Page(t("campaign.mission_deck.title.situation"), blocks)


def _ato_page(picture: MissionPicture) -> Page:
    package = []
    for flight in picture.flights:
        package.append(
            t(
                "campaign.mission_deck.flight",
                callsign=flight.callsign,
                name=flight.name,
                count=flight.count,
                aircraft=flight.aircraft,
            )
        )
        package.append(
            t("campaign.mission_deck.flight_base", base=flight.base)
            if flight.base
            else t("campaign.mission_deck.flight_air_start")
        )
        package.append(
            "  ".join(f"{n}. ______" for n in range(1, flight.count + 1))
            + "    "
            + t("campaign.mission_deck.free_loadout")
        )
    for airfield in picture.airfields:
        package.append(t("campaign.mission_deck.dynamic_slots", airfield=airfield.name))
        package.append("1. ______  2. ______  3. ______  4. ______    " + t("campaign.mission_deck.free_loadout"))
    blocks = [Block(t("campaign.mission_deck.heading.package"), tuple(package))] if package else []
    support = []
    for aircraft in picture.support:
        tacan = f", TACAN {aircraft.tacan}" if aircraft.tacan else ""
        support.append(
            t(
                "campaign.mission_deck.support",
                callsign=aircraft.callsign,
                role=t(f"campaign.mission_deck.role.{aircraft.role}"),
                aircraft=aircraft.aircraft,
                frequency=_mhz(aircraft.frequency),
                tacan=tacan,
            )
        )
    if support:
        blocks.append(Block(t("campaign.mission_deck.heading.support"), tuple(support)))
    control = [_carrier_line(carrier) for carrier in picture.carriers]
    control += [
        t("campaign.mission_deck.airfield", airfield=a.name, uhf=_mhz(a.uhf), vhf=_mhz(a.vhf), tacan=a.tacan or "—")
        for a in picture.airfields
    ]
    if control:
        blocks.append(Block(t("campaign.mission_deck.heading.control"), tuple(control)))
    return Page(t("campaign.mission_deck.title.ato"), blocks)


def _carrier_line(carrier: Carrier) -> str:
    parts = [f"{_mhz(carrier.tower)} MHz"]
    if carrier.tacan:
        parts.append(f"TACAN {carrier.tacan}")
    if carrier.icls is not None:
        parts.append(f"ICLS {carrier.icls}")
    if carrier.link4 is not None:
        parts.append(f"Link 4 {_mhz(carrier.link4)} MHz")
    return t("campaign.mission_deck.carrier", carrier=carrier.name, channels=", ".join(parts))


def _flow_page(
    campaign: CampaignDefinition, state: CampaignState, page: MissionPage | None, picture: MissionPicture
) -> Page:
    blocks = []
    if page and page.tasks:
        blocks.append(
            Block(
                t("campaign.mission_deck.heading.objectives"),
                tuple(f"{n}. {task.title}" for n, task in enumerate(page.tasks, 1)),
            )
        )
    qra = _qra_line(campaign, picture)
    blocks.append(Block(t("campaign.mission_deck.heading.air"), (qra or t("campaign.mission_deck.no_air"),)))
    enemy = enemy_of(campaign.player_side)
    defences = [
        _line(zone.label, _after_colon(enemy_picture(campaign, state, zone)))
        for zone in campaign.zones
        if state.zones[zone.name].owner == enemy and campaign.size_classes[zone.size].long_range_sam
    ]
    defences.append(t("campaign.mission_deck.escort_defences"))
    blocks.append(Block(t("campaign.mission_deck.heading.air_defence"), tuple(defences)))
    others = []
    tankers = [s.callsign for s in picture.support if s.role != "awacs"]
    if tankers:
        others.append(t("campaign.mission_deck.refuelling", tankers=_join(tankers)))
    diversions = [a.name for a in picture.airfields] + [c.name for c in picture.carriers]
    if diversions:
        others.append(t("campaign.mission_deck.diversion", fields=_join(diversions)))
    if picture.plane_guard:
        others.append(t("campaign.mission_deck.plane_guard"))
    if picture.csar:
        others.append(t("campaign.mission_deck.csar"))
    if others:
        blocks.append(Block(t("campaign.mission_deck.heading.other"), tuple(others)))
    return Page(t("campaign.mission_deck.title.flow"), blocks)


def _frequencies_page(picture: MissionPicture) -> Page:
    uhf = [t("campaign.mission_deck.guard", frequency=_mhz(_GUARD))]
    uhf += [
        _line(s.callsign, f"{_mhz(s.frequency)} MHz" + (f" — TACAN {s.tacan}" if s.tacan else ""))
        for s in picture.support
    ]
    uhf += [_line(a.name, f"{_mhz(a.uhf)} MHz" + (f" — TACAN {a.tacan}" if a.tacan else "")) for a in picture.airfields]
    vhf = [_carrier_line(carrier) for carrier in picture.carriers]
    vhf += [
        _line(a.name, f"{_mhz(a.vhf)} MHz" + (f" (FM {_mhz(a.fm)} MHz)" if a.fm else "")) for a in picture.airfields
    ]
    blocks = [Block("UHF", tuple(uhf))]
    if vhf:
        blocks.append(Block("VHF", tuple(vhf)))
    return Page(t("campaign.mission_deck.title.frequencies"), blocks)


def _coordinates_page(campaign: CampaignDefinition, zones: list[CampaignZone]) -> Page:
    grid = grid_for_theatre(campaign.theatre)
    lines = [t("campaign.mission_deck.coordinates_format")]
    for zone in zones:
        lat, lon = zone_position(campaign, zone)
        height = grid.elevation_at(*_zone_xy(campaign, zone)) if grid else None
        # no swept terrain on this workstation: nothing rather than "unknown"
        elevation = f" — {metres_to_feet(height):.0f} ft" if height is not None else ""
        lines.append(_line(zone.label, f"{to_dms(lat, lon)}{elevation}"))
    lines.append(t("campaign.mission_deck.positions_in_flight"))
    return Page(t("campaign.mission_deck.title.coordinates"), [Block(None, tuple(lines))])


def _picture_page(deck: PresentationType, title: str, image: Path, notes: list[tuple[str, bool]]) -> None:
    """A map on the left, its notes on the right; a note is ``(text, is_heading)``."""
    slide = _titled_slide(deck, title)
    with Image.open(image) as opened:
        ratio = opened.width / opened.height
    height = _HEIGHT - Emu(1250000)
    width = min(int(height * ratio), Emu(8000000))
    height = int(width / ratio)
    slide.shapes.add_picture(str(image), _MARGIN, Emu(1050000), width=Emu(width), height=Emu(height))
    frame = _frame(slide, _MARGIN + width + Emu(250000), Emu(1150000), _WIDTH - width - 3 * _MARGIN, Emu(5000000))
    for index, (text, heading) in enumerate(notes):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _run(paragraph, text, 15 if heading else 13, bold=heading)


def _write_cover(
    deck: PresentationType,
    campaign: CampaignDefinition,
    prose: BriefingProse | None,
    page: MissionPage | None,
    picture: MissionPicture,
    mission: int,
    contents: list[str],
) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = _frame(slide, _MARGIN, Emu(1700000), Emu(6800000), Emu(3200000))
    operation = (prose.operation if prose and prose.operation else campaign.name).upper()
    _run(frame.paragraphs[0], f"{operation} {mission}/{campaign.missions}", 40, bold=True)
    _run(
        frame.add_paragraph(),
        page.title if page else t("campaign.mission_deck.untitled", mission=mission),
        54,
        bold=True,
        colour=_BLUE,
    )
    _run(frame.add_paragraph(), t("campaign.mission_deck.subtitle"), 22, colour=_GREY)
    date = t("campaign.mission_deck.date", day=picture.date.day, month=picture.date.month, year=picture.date.year)
    _run(
        frame.add_paragraph(),
        f"{date} — {t('campaign.mission_deck.local', time=_clock(picture.start_time))}",
        16,
        italic=True,
        colour=_GREY,
    )
    toc = _frame(slide, Emu(7900000), Emu(1500000), Emu(3900000), Emu(4600000))
    _run(toc.paragraphs[0], t("campaign.deck.contents"), 20, bold=True)
    for item in contents:
        _run(toc.add_paragraph(), item, 14)


def mission_deck(
    campaign: CampaignDefinition,
    state: CampaignState,
    prose: BriefingProse | None,
    picture: MissionPicture,
    folder: Path,
    *,
    cache_dir: Path | None = None,
    fetch: Fetch | None = None,
) -> MissionDeckReport:
    """Write the coming mission's briefing, and its maps, into the mission's sub-folder.

    Args:
        campaign: The validated campaign.
        state: The campaign state the coming mission starts from.
        prose: `briefing.yaml`, for the mission's title and tasks, or ``None``.
        picture: What the built mission says.
        folder: The mission's sub-folder, `missions/mission-NN/`.
        cache_dir: Where map tiles are kept between runs.
        fetch: How a map tile is downloaded (tests pass their own).

    Returns:
        Where the deck is, and what it was made with.
    """
    mission = state.mission + 1
    page = prose.missions.get(mission) if prose else None
    zones = objective_zones(campaign, page)
    folder.mkdir(parents=True, exist_ok=True)
    tactical = render_tactical_map(
        campaign, state, picture, folder / TACTICAL_MAP_FILE, cache_dir=cache_dir, fetch=fetch
    )
    zooms = [
        render_objective_map(campaign, state, zone, folder / zoom_file(zone), cache_dir=cache_dir, fetch=fetch)
        for zone in zones
    ]

    situation = _situation_page(campaign, state, prose, page, picture, mission)
    ato = _ato_page(picture)
    flow = _flow_page(campaign, state, page, picture)
    frequencies = _frequencies_page(picture)
    coordinates = _coordinates_page(campaign, zones)
    titles = [
        situation.title,
        ato.title,
        t("campaign.mission_deck.title.tactical"),
        *[t("campaign.mission_deck.title.zoom", zone=zone.label) for zone in zones],
        flow.title,
        frequencies.title,
        coordinates.title,
    ]
    numbered = [f"{n}. {title}" for n, title in enumerate(titles, 1)]
    situation.title, ato.title = numbered[0], numbered[1]
    flow.title, frequencies.title, coordinates.title = numbered[-3:]

    deck = Presentation()
    deck.slide_width, deck.slide_height = _WIDTH, _HEIGHT
    _write_cover(deck, campaign, prose, page, picture, mission, numbered)
    for part in [*paginate(situation), *paginate(ato)]:
        _write_page(deck, part)
    legend = [
        (t("campaign.deck.legend"), True),
        (t("campaign.mission_deck.legend.zones"), False),
        (t("campaign.mission_deck.legend.qra"), False),
        (t("campaign.mission_deck.legend.carrier"), False),
        (t("campaign.mission_deck.legend.support"), False),
        (t("campaign.mission_deck.legend.remember"), True),
        (t("campaign.mission_deck.positions_in_flight"), False),
    ]
    _picture_page(deck, numbered[2], tactical.path, legend)
    for title, zone, zoom in zip(numbered[3 : 3 + len(zones)], zones, zooms, strict=True):
        notes = [
            (zone.label, True),
            (_status(campaign, state, zone), False),
            (t("campaign.mission_deck.zone_radius", km=f"{zone.radius / 1000:g}"), False),
        ]
        task = next((task for task in (page.tasks if page else ()) if _names(task.title, zone)), None)
        if task:
            notes += [(t("campaign.mission_deck.heading.task"), True), (task.title, False)]
            notes += [(text, False) for text in task.text]
        _picture_page(deck, title, zoom.path, notes)
    for part in [*paginate(flow), *paginate(frequencies), *paginate(coordinates)]:
        _write_page(deck, part)
    out = folder / MISSION_DECK_FILE
    deck.save(str(out))
    return MissionDeckReport(
        path=out,
        pages=len(deck.slides),
        maps=[tactical.path, *(zoom.path for zoom in zooms)],
        offline=tactical.offline or any(zoom.offline for zoom in zooms),
    )
