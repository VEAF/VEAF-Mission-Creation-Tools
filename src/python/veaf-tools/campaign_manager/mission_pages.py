"""The mission briefing's pages that only read the built mission (FEAT-MCP-MISSION-BRIEFING ticket 01).

Shared by the campaign mission's deck (`mission_deck.py`) and a one-shot mission's
(`mission_briefing/`): the ATO, the frequency plan, the flight plan, the weather, and a map page.
Nothing here knows a campaign.
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx.presentation import Presentation as PresentationType
from pptx.util import Emu
from veaf_libs.coordinates import to_dms, xy_to_latlon
from veaf_libs.i18n import t
from veaf_libs.terrain_elevation import metres_to_feet

from campaign_manager.briefing_deck import _HEIGHT, _MARGIN, _WIDTH, Block, Page, _frame, _run, _titled_slide
from campaign_manager.mission_picture import Carrier, MissionPicture, Waypoint

#: hPa per mmHg.
_HPA = 1.33322


def line(name: str, text: str) -> str:
    """``name: text``, punctuated as the language wants it."""
    return t("campaign.mission_deck.line", name=name, text=text)


def clock(seconds: int) -> str:
    """``14h30`` from seconds since midnight."""
    return f"{seconds // 3600:02d}h{seconds % 3600 // 60:02d}"


def mhz(value: float | None) -> str:
    """A frequency in MHz, ``?`` when unknown."""
    return "?" if value is None else f"{value:.3f}".rstrip("0").rstrip(".") if value % 1 else f"{value:.1f}"


def weather_lines(picture: MissionPicture) -> tuple[str, ...]:
    """The weather as the briefing says it: sky, visibility, temperature, QNH, then the winds."""
    weather = picture.weather
    sky = (
        t("campaign.mission_deck.sky_clear")
        if weather.cover == "clear"
        else t(
            "campaign.mission_deck.sky_clouds",
            cover=t(f"campaign.mission_deck.cover.{weather.cover}"),
            base=round(metres_to_feet(weather.cloud_base), -2),
        )
    )
    visibility = (
        t("campaign.mission_deck.visibility_10")
        if weather.visibility >= 9000
        else t("campaign.mission_deck.visibility_km", km=round(weather.visibility / 1000))
    )
    extras = [t("campaign.mission_deck.rain")] if weather.rain else []
    extras += [t("campaign.mission_deck.fog")] if weather.fog else []
    winds = weather.winds
    return (
        t(
            "campaign.mission_deck.weather",
            sky=sky,
            visibility=visibility,
            temperature=round(weather.temperature),
            qnh=round(weather.qnh),
            hpa=round(weather.qnh * _HPA),
        ),
        t(
            "campaign.mission_deck.wind",
            ground=f"{winds['ground'].origin:03d}° / {winds['ground'].knots} kt",
            at2000=f"{winds['2000'].origin:03d}° / {winds['2000'].knots} kt",
            at8000=f"{winds['8000'].origin:03d}° / {winds['8000'].knots} kt",
        ),
        *extras,
    )


def carrier_line(carrier: Carrier) -> str:
    """A carrier and its channels: tower, TACAN, ICLS, Link 4."""
    parts = [f"{mhz(carrier.tower)} MHz"]
    if carrier.tacan:
        parts.append(f"TACAN {carrier.tacan}")
    if carrier.icls is not None:
        parts.append(f"ICLS {carrier.icls}")
    if carrier.link4 is not None:
        parts.append(f"Link 4 {mhz(carrier.link4)} MHz")
    return t("campaign.mission_deck.carrier", carrier=carrier.name, channels=", ".join(parts))


def ato_page(
    picture: MissionPicture, pilots: dict[str, list[str]] | None = None, loadouts: dict[str, str] | None = None
) -> Page:
    """The package, the support assets and the control.

    Args:
        picture: What the built mission says.
        pilots: The pilots' names by flight group name; empty cells otherwise.
        loadouts: The imposed loadout by flight group name; "free loadout" otherwise.

    Returns:
        The page.
    """
    pilots, loadouts = pilots or {}, loadouts or {}
    package = []
    for flight in picture.flights:
        package.append(
            t(
                "campaign.mission_deck.flight",
                callsign=flight.callsign,
                name=flight.name,
                count=flight.count,
                aircraft=flight.aircraft,
            )
        )
        package.append(
            t("campaign.mission_deck.flight_base", base=flight.base)
            if flight.base
            else t("campaign.mission_deck.flight_air_start")
        )
        names = pilots.get(flight.name, [])
        package.append(
            "  ".join(f"{n}. {names[n - 1] if n <= len(names) else '______'}" for n in range(1, flight.count + 1))
            + "    "
            + loadouts.get(flight.name, t("campaign.mission_deck.free_loadout"))
        )
    for airfield in picture.airfields:
        package.append(t("campaign.mission_deck.dynamic_slots", airfield=airfield.name))
        package.append("1. ______  2. ______  3. ______  4. ______    " + t("campaign.mission_deck.free_loadout"))
    blocks = [Block(t("campaign.mission_deck.heading.package"), tuple(package))] if package else []
    support = []
    for aircraft in picture.support:
        tacan = f", TACAN {aircraft.tacan}" if aircraft.tacan else ""
        support.append(
            t(
                "campaign.mission_deck.support",
                callsign=aircraft.callsign,
                role=t(f"campaign.mission_deck.role.{aircraft.role}"),
                aircraft=aircraft.aircraft,
                frequency=mhz(aircraft.frequency),
                tacan=tacan,
            )
        )
    if support:
        blocks.append(Block(t("campaign.mission_deck.heading.support"), tuple(support)))
    control = [carrier_line(carrier) for carrier in picture.carriers]
    control += [
        t("campaign.mission_deck.airfield", airfield=a.name, uhf=mhz(a.uhf), vhf=mhz(a.vhf), tacan=a.tacan or "—")
        for a in picture.airfields
    ]
    if control:
        blocks.append(Block(t("campaign.mission_deck.heading.control"), tuple(control)))
    return Page(t("campaign.mission_deck.title.ato"), blocks)


def frequencies_page(picture: MissionPicture) -> Page:
    """UHF then VHF: the support, the airfields, the carriers. No Guard: every aircraft reaches it by its own control."""
    uhf = [
        line(s.callsign, f"{mhz(s.frequency)} MHz" + (f" — TACAN {s.tacan}" if s.tacan else ""))
        for s in picture.support
    ]
    uhf += [line(a.name, f"{mhz(a.uhf)} MHz" + (f" — TACAN {a.tacan}" if a.tacan else "")) for a in picture.airfields]
    vhf = [carrier_line(carrier) for carrier in picture.carriers]
    vhf += [line(a.name, f"{mhz(a.vhf)} MHz" + (f" (FM {mhz(a.fm)} MHz)" if a.fm else "")) for a in picture.airfields]
    blocks = [Block("UHF", tuple(uhf))] if uhf else []
    if vhf:
        blocks.append(Block("VHF", tuple(vhf)))
    return Page(t("campaign.mission_deck.title.frequencies"), blocks)


def altitude(waypoint: Waypoint) -> str:
    """``10 000 ft BARO``, ``500 ft AGL`` or ``ground``, to the ten feet, never broken across two lines."""
    if waypoint.alt_type == "RADIO" and waypoint.alt == 0:
        return t("campaign.mission_deck.altitude.ground")
    feet = f"{round(metres_to_feet(waypoint.alt), -1):,.0f}".replace(",", " ")
    return t(f"campaign.mission_deck.altitude.{'agl' if waypoint.alt_type == 'RADIO' else 'baro'}", feet=feet)


def navigation_page(theatre: str, picture: MissionPicture) -> Page | None:
    """The players' waypoints after their departure, one block per route: planes, then helicopters.

    Numbered as the mission editor numbers them, the departure being 0; a category carrying several
    routes names the types that fly each.
    """
    blocks = []
    for plan in picture.flight_plans:
        heading = t(f"campaign.mission_deck.heading.{plan.category}s")
        if sum(1 for other in picture.flight_plans if other.category == plan.category) > 1:
            heading = f"{heading} — {', '.join(plan.aircraft)}"
        lines = tuple(
            f"{index}. " + line(wp.name, f"{to_dms(*xy_to_latlon(theatre, wp.x, wp.y))} — {altitude(wp)}")
            for index, wp in enumerate(plan.waypoints, 1)
        )
        blocks.append(Block(heading, lines))
    return Page(t("campaign.mission_deck.title.navigation"), blocks) if blocks else None


def picture_page(deck: PresentationType, title: str, image: Path, notes: list[tuple[str, bool]]) -> None:
    """A map on the left, its notes on the right; a note is ``(text, is_heading)``."""
    slide = _titled_slide(deck, title)
    with Image.open(image) as opened:
        ratio = opened.width / opened.height
    height = _HEIGHT - Emu(1250000)
    width = min(int(height * ratio), Emu(8000000))
    height = int(width / ratio)
    slide.shapes.add_picture(str(image), _MARGIN, Emu(1050000), width=Emu(width), height=Emu(height))
    frame = _frame(slide, _MARGIN + width + Emu(250000), Emu(1150000), _WIDTH - width - 3 * _MARGIN, Emu(5000000))
    for index, (text, heading) in enumerate(notes):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        _run(paragraph, text, 15 if heading else 13, bold=heading)
