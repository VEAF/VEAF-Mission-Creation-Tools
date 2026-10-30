"""Sheets of cards: the targets and the flight plan's steps, six to a slide (FEAT-OBJECTIVE-BRIEFING-QUALITY 03, 05).

A list of names and coordinates is what the briefing had for its targets, and David found it arid: a
pilot looks for one target, its coordinates in the format his aircraft takes, its elevation, and where
it stands. A card gives each one its box: a thumbnail cut from the map that shows it, its name, and its
lines. The flight plan's steps are cards too — the point, the heading and distance that reach it, the
highest ground on the leg, what is done there —, as a DCS kneeboard shows a route one point at a time.

The deck draws them with python-pptx shapes in the template's font (Google Slides knows Calibri), the PDF
with fpdf2, from the same :class:`CardPage`.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path

from campaign_manager.briefing_deck import _HEIGHT, _MARGIN, _WIDTH, _frame, _run, _titled_slide
from campaign_manager.strategic_map import OWNER_COLOURS
from PIL import Image, ImageDraw
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.presentation import Presentation as PresentationType
from pptx.util import Emu, Pt
from veaf_libs.i18n import t

#: Cards on a slide: two columns, three rows.
COLUMNS, ROWS = 2, 3
PER_PAGE = COLUMNS * ROWS

#: A thumbnail's side, pixels, before it is fitted into its card.
THUMBNAIL = 360

_BORDER = RGBColor(0xB0, 0xB0, 0xB0)
_SUBTITLE = RGBColor(0x1F, 0x5F, 0xBF)


@dataclass(frozen=True)
class Card:
    """One card: a title, its lines, and the thumbnail beside them (JPEG), if any."""

    title: str
    lines: tuple[str, ...]
    thumbnail: bytes | None = None


@dataclass
class CardPage:
    """A slide of cards under its title, and the line that says what they are (an objective, a route)."""

    title: str
    heading: str | None
    cards: list[Card]


@dataclass
class CardSection:
    """A part of the briefing made of cards: one title in the contents, its groups paginated six by six."""

    title: str
    groups: list[tuple[str | None, list[Card]]] = field(default_factory=list)

    def pages(self) -> list[CardPage]:
        """The section's slides: a group never shares one, a slide after the first says it continues."""
        pages: list[CardPage] = []
        for heading, cards in self.groups:
            for start in range(0, len(cards), PER_PAGE):
                title = self.title if not pages else t("campaign.deck.continued", title=self.title)
                pages.append(CardPage(title, heading, cards[start : start + PER_PAGE]))
        return pages


def thumbnail(image: Path, centre: tuple[float, float], side: int = THUMBNAIL) -> bytes:
    """A square cut from a map around a point, the point ringed: what a card shows.

    Args:
        image: The rendered map.
        centre: The point, in the map's pixels.
        side: The square's side, pixels; smaller when the map is.

    Returns:
        The square, JPEG: sixty of them in PNG took the deck to 11 MB.
    """
    with Image.open(image) as opened:
        picture = opened.convert("RGB")
    side = min(side, picture.width, picture.height)
    left = min(max(round(centre[0] - side / 2), 0), picture.width - side)
    top = min(max(round(centre[1] - side / 2), 0), picture.height - side)
    square = picture.crop((left, top, left + side, top + side))
    cx, cy = centre[0] - left, centre[1] - top
    radius = side / 9
    ImageDraw.Draw(square).ellipse(
        (cx - radius, cy - radius, cx + radius, cy + radius), outline=OWNER_COLOURS["red"], width=max(2, side // 90)
    )
    buffer = io.BytesIO()
    square.save(buffer, format="JPEG", quality=85)
    return buffer.getvalue()


# the slide's geometry, EMU
_GAP = Emu(150000)
_PAD = Emu(70000)
_HEADING_TOP = Emu(1000000)
_CARDS_TOP = Emu(1420000)
_CARDS_BOTTOM = _HEIGHT - Emu(230000)


def card_boxes() -> list[tuple[int, int, int, int]]:
    """Each card's ``(left, top, width, height)`` on the slide, row by row."""
    width = (_WIDTH - 2 * _MARGIN - (COLUMNS - 1) * _GAP) // COLUMNS
    height = (_CARDS_BOTTOM - _CARDS_TOP - (ROWS - 1) * _GAP) // ROWS
    return [
        (_MARGIN + column * (width + _GAP), _CARDS_TOP + row * (height + _GAP), width, height)
        for row in range(ROWS)
        for column in range(COLUMNS)
    ]


def write_card_page(deck: PresentationType, page: CardPage) -> None:
    """A slide of cards: its title, the line saying what they are, then each card in its box."""
    slide = _titled_slide(deck, page.title)
    if page.heading:
        frame = _frame(slide, _MARGIN, _HEADING_TOP, _WIDTH - 2 * _MARGIN, Emu(380000))
        _run(frame.paragraphs[0], page.heading, 16, bold=True, colour=_SUBTITLE)
    for card, (left, top, width, height) in zip(page.cards, card_boxes(), strict=False):
        box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(left), Emu(top), Emu(width), Emu(height))
        box.fill.background()
        box.line.color.rgb = _BORDER
        box.line.width = Pt(0.75)
        text_left = left + _PAD
        if card.thumbnail:
            side = height - 2 * _PAD
            slide.shapes.add_picture(
                io.BytesIO(card.thumbnail), Emu(left + _PAD), Emu(top + _PAD), width=Emu(side), height=Emu(side)
            )
            text_left += side + _PAD
        frame = _frame(slide, text_left, top + _PAD, left + width - _PAD - text_left, height - 2 * _PAD)
        _run(frame.paragraphs[0], card.title, 14, bold=True)
        for text in card.lines:
            _run(frame.add_paragraph(), text, 11)
