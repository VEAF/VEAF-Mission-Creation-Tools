"""Pages of tables: the ATO, the deployment and the code words (FIX-BRIEFING-PARITY-MARTEAU 07).

The hand-made briefing of "Marteau du Golan" set them as tables, a column for when, one for who, one for
what; the tool printed them as plain lines, the times lost in the prose. A :class:`TablePage` is one or
more tables under a title, and the prose that follows them; the deck draws it with python-pptx's tables,
the PDF with fpdf2's, from the same rows.
"""

from __future__ import annotations

import math
import re
from dataclasses import dataclass, field

from campaign_manager.briefing_deck import _HEIGHT, _MARGIN, _WIDTH, Block, _frame, _run, _titled_slide
from pptx.dml.color import RGBColor
from pptx.presentation import Presentation as PresentationType
from pptx.table import _Cell
from pptx.util import Emu, Pt

#: The deck's text sizes in a table, points.
_HEADER_SIZE, _CELL_SIZE = 12, 11

#: The room a table page has for its rows, in lines of :data:`_CELL_SIZE` text.
_PAGE_LINES = 26

#: How many characters of :data:`_CELL_SIZE` text a full-width line holds, on the deck's body.
_LINE_CHARS = 160

_TOP = Emu(1050000)
_LINE_EMU = Emu(190000)
_HEADER_FILL = RGBColor(0x1F, 0x4E, 0xA8)
_STRIPE_FILL = RGBColor(0xF2, 0xF4, 0xF8)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)


@dataclass(frozen=True)
class Table:
    """A table: its heading (or none), its columns' titles and shares of the width, its rows."""

    heading: str | None
    columns: tuple[str, ...]
    widths: tuple[float, ...]
    """Each column's share of the width; they sum to 1."""
    rows: tuple[tuple[str, ...], ...]


@dataclass
class TablePage:
    """A page of tables under its title, then its prose."""

    title: str
    tables: list[Table]
    after: list[Block] = field(default_factory=list)


def row_lines(table: Table, row: tuple[str, ...]) -> int:
    """How many lines a row takes: its tallest cell, its text wrapped in its column."""
    tallest = 1
    for text, share in zip(row, table.widths, strict=True):
        room = max(8, int(_LINE_CHARS * share) - 2)
        tallest = max(tallest, sum(max(1, math.ceil(len(part) / room)) for part in text.split("\n")))
    return tallest


def _block_lines(block: Block) -> int:
    lines = 2 if block.heading else 0
    return lines + sum(max(1, math.ceil(len(text) / _LINE_CHARS)) for text in block.paragraphs)


def paginate_tables(page: TablePage) -> list[TablePage]:
    """Cut a page of tables into slides: a table runs on with its header repeated, the prose follows.

    Args:
        page: The page, however long.

    Returns:
        Its slides, the first under the page's title, the others under the same title (the caller numbers
        or marks them).
    """
    pages = [TablePage(page.title, [])]
    used = 0
    for table in page.tables:
        rows: list[tuple[str, ...]] = []
        header = 2 + (2 if table.heading else 0)
        for row in table.rows:
            lines = row_lines(table, row)
            if rows and used + header + sum(row_lines(table, r) for r in rows) + lines > _PAGE_LINES:
                pages[-1].tables.append(Table(table.heading, table.columns, table.widths, tuple(rows)))
                pages.append(TablePage(page.title, []))
                used, rows = 0, []
            rows.append(row)
        if rows:
            pages[-1].tables.append(Table(table.heading, table.columns, table.widths, tuple(rows)))
            used += header + sum(row_lines(table, r) for r in rows) + 1
    for block in page.after:
        lines = _block_lines(block)
        if used and used + lines > _PAGE_LINES:
            pages.append(TablePage(page.title, []))
            used = 0
        pages[-1].after.append(block)
        used += lines
    return pages


#: A deployment line: ``H+15 — Archer, Jedi : décollage…`` — when, who, what.
_TIMELINE = re.compile(r"^(?P<when>[^—:]{1,40}?)\s+—\s+(?P<rest>.+)$")

#: Who acts, before the colon: a few names, never a sentence.
_WHO = re.compile(r"^(?P<who>[^:.]{1,40}?)\s*:\s+(?P<what>.+)$")


def timeline(paragraphs: tuple[str, ...]) -> tuple[list[tuple[str, str, str]], list[str]]:
    """Split the deployment's paragraphs into the timeline's rows and the prose around it.

    A row is a paragraph ``<when> — <who> : <what>``; a paragraph whose part before the colon reads as a
    sentence keeps it in ``what`` and leaves ``who`` empty.

    Args:
        paragraphs: The deployment, as the assistant wrote it.

    Returns:
        ``(rows, prose)``: each row ``(when, who, what)``, then the paragraphs that are no row.
    """
    rows: list[tuple[str, str, str]] = []
    prose: list[str] = []
    for text in paragraphs:
        found = _TIMELINE.match(text)
        if found is None:
            prose.append(text)
            continue
        rest = found.group("rest")
        who = _WHO.match(rest)
        if who and len(who.group("who").split()) <= 6:
            rows.append((found.group("when"), who.group("who"), who.group("what")))
        else:
            rows.append((found.group("when"), "", rest))
    return rows, prose


def _cell(cell: _Cell, text: str, size: int, *, bold: bool = False, colour: RGBColor | None = None) -> None:
    frame = cell.text_frame
    frame.word_wrap = True
    _run(frame.paragraphs[0], text, size, bold=bold, colour=colour)
    cell.margin_left = cell.margin_right = Emu(60000)
    cell.margin_top = cell.margin_bottom = Emu(30000)


def write_table_page(deck: PresentationType, page: TablePage) -> None:
    """A slide of tables: each under its heading, header row filled, rows striped; the prose under them."""
    slide = _titled_slide(deck, page.title)
    width = _WIDTH - 2 * _MARGIN
    top = int(_TOP)
    for table in page.tables:
        if table.heading:
            _run(_frame(slide, _MARGIN, top, width, Emu(330000)).paragraphs[0], table.heading, 16, bold=True)
            top += int(Emu(380000))
        heights = [int(_LINE_EMU) + int(Emu(60000))] + [
            row_lines(table, row) * int(_LINE_EMU) + int(Emu(60000)) for row in table.rows
        ]
        total = sum(heights, 0)
        shape = slide.shapes.add_table(
            len(table.rows) + 1, len(table.columns), Emu(_MARGIN), Emu(top), Emu(width), Emu(total)
        )
        grid = shape.table
        for index, share in enumerate(table.widths):
            grid.columns[index].width = Emu(int(width * share))
        for index, height in enumerate(heights):
            grid.rows[index].height = Emu(height)
        for column, title in enumerate(table.columns):
            cell = grid.cell(0, column)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _HEADER_FILL
            _cell(cell, title, _HEADER_SIZE, bold=True, colour=_WHITE)
        for number, row in enumerate(table.rows, 1):
            for column, text in enumerate(row):
                cell = grid.cell(number, column)
                cell.fill.solid()
                cell.fill.fore_color.rgb = _STRIPE_FILL if number % 2 == 0 else _WHITE
                _cell(cell, text, _CELL_SIZE, bold=column == 0)
        top += total + int(Emu(200000))
    if page.after:
        frame = _frame(slide, _MARGIN, top, width, max(int(Emu(400000)), int(_HEIGHT) - top - int(Emu(300000))))
        first = True
        for block in page.after:
            if block.heading:
                paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
                first = False
                paragraph.space_before = Pt(6)
                _run(paragraph, block.heading, 14, bold=True)
            for text in block.paragraphs:
                paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
                first = False
                _run(paragraph, text, 12)
