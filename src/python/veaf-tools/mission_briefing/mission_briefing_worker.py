"""Build a one-shot mission's VEAF briefing from the built mission (FEAT-MCP-MISSION-BRIEFING ticket 04).

The objective-mission prompt's §3, page for page, in the mission's language: what the assistant wrote
(title, context, mission, objectives, threats, flow) around what the **built** `.miz` holds — date,
time, weather, bullseye, flights, support, carriers, QRA, flight plan, the presets in the cockpits,
each target where it stands. Then the maps become the DCS briefing pictures, in the mission folder
(durable) and in the built `.miz` (no rebuild needed), listed first for the players' side and neutral; the
other side is left alone (`briefing-pictures-red-then-blue`). The pictures this writes carry :data:`PICTURE_PREFIX` keys
and are replaced on every run.
"""

from __future__ import annotations

import io
import math
import re
import shutil
import tempfile
import unicodedata
from dataclasses import replace
from datetime import datetime
from pathlib import Path
from typing import Any

import luadata
from campaign_manager.briefing_deck import (
    _BLUE,
    _GREY,
    _HEIGHT,
    _MARGIN,
    _WIDTH,
    Block,
    Page,
    _frame,
    _join,
    _run,
    _titled_slide,
    _write_page,
    paginate,
)
from campaign_manager.map_canvas import TacticalMapReport
from campaign_manager.mission_pages import (
    altitude,
    carrier_line,
    clock,
    frequencies_page,
    line,
    mhz,
    navigation_page,
    picture_page,
    weather_lines,
)
from campaign_manager.mission_picture import (
    NM,
    FlightPlan,
    MissionPicture,
    _flight_plans,
    find_built_mission,
    read_mission_picture,
)
from mission_tools.mission_constants import DEFAULT_SCRIPTS_LOCATION
from mission_tools.mission_yaml_editor import load_yaml
from mission_tools.miz_backup import backup_before_write
from mission_tools.miz_tools import DcsMission, list_members, read_miz, write_miz
from PIL import Image
from pptx import Presentation
from pptx.presentation import Presentation as PresentationType
from pptx.util import Emu
from presets_injector.presets_manager import KNEEBOARD_CALIBRATION_PATH
from veaf_libs.coordinates import to_dms, xy_to_latlon
from veaf_libs.i18n import current_language, language, t
from veaf_libs.map_tiles import OSM, Fetch, TileSource
from veaf_libs.mission_table import indexed
from veaf_libs.terrain_elevation import grid_for_theatre, metres_to_feet
from veaf_mission_mcp.add_sound import _folder_l10n
from veaf_mission_mcp.mission_folder import load_folder_mission, save_folder_mission

from mission_briefing.briefing_cards import Card, CardPage, CardSection, thumbnail, to_ddm, write_card_page
from mission_briefing.briefing_kneeboard import (
    KNEEBOARD_FOLDER,
    MAP_ASPECT,
    flight_card_blocks,
    is_kneeboard_page,
    leg_text,
    map_page,
    page_name,
    photo_page,
    route_steps,
    targets_blocks,
    text_pages,
)
from mission_briefing.briefing_maps import (
    COMBAT_AREA_FILE,
    OVERVIEW_FILE,
    Inset,
    attack_axis_rotation,
    render_objective_zoom,
    render_overview_map,
    zoom_file,
    zoom_half,
)
from mission_briefing.briefing_pdf import Cover, MapSheet, PhotoSheet, write_pdf
from mission_briefing.briefing_recon import RECON_FOLDER, caption, load_recon
from mission_briefing.briefing_tables import Table, TablePage, paginate_tables, timeline, write_table_page
from mission_briefing.mission_reader import (
    departures,
    ground_picture,
    group_position,
    objective_area,
    operation_zones,
    radio_plans,
)
from mission_briefing.models import (
    BriefingContent,
    BriefingReport,
    GroundPicture,
    Objective,
    ObjectiveArea,
    PackageFlight,
    RadioPlan,
    ReconPhoto,
    Threat,
)

#: The deck, in the mission folder.
DECK_FILE = Path("docs") / "briefing.pptx"

#: The PDF of the same pages, in the mission folder.
PDF_FILE = Path("docs") / "briefing.pdf"

#: The maps, in the mission folder.
MAPS_FOLDER = Path("docs") / "maps"

#: The folder the build writes the weather variants into, beside its own `.miz`.
VARIANTS_FOLDER = "missions"

#: The resource keys of the pictures this writes; anything else in the briefing is left alone.
PICTURE_PREFIX = "MCP_Briefing_"

#: The width a DCS briefing picture is brought down to, pixels.
_PICTURE_WIDTH = 1600

#: The overview's width, pixels, and its aspect: the slide's body, under the title, the whole width
#: (FEAT-OBJECTIVE-BRIEFING-QUALITY 02: 715 px beside a legend taking half the slide could not be read).
_OVERVIEW = 2000
_OVERVIEW_ASPECT = 2.0

#: An objective zoom's width, pixels.
_ZOOM = 1600

#: A kneeboard map's width, pixels: the page's, inside its margins.
_KNEEBOARD_MAP = 1392

#: The mission-table key of each side's briefing pictures; the slots of no side get them too, as ``neutral``.
_PICTURE_KEYS = {"blue": "pictureFileNameB", "red": "pictureFileNameR", "neutral": "pictureFileNameN"}


def _mission_language(folder: Path) -> str:
    """`mission.language` of `mission.yaml`, the tools' language otherwise."""
    path = folder / "mission.yaml"
    mission = (load_yaml(path).get("mission") if path.is_file() else None) or {}
    return str(mission.get("language") or current_language()) if isinstance(mission, dict) else current_language()


def _bullseye_line(theatre: str, picture: MissionPicture, start: tuple[str, float, float] | None) -> str:
    """The bullseye in DMS, and its bearing and range from the players' departure when there is one."""
    lat, lon = xy_to_latlon(theatre, *picture.bullseye)
    if start is None:
        return to_dms(lat, lon)
    name, x, y = start
    dx, dy = picture.bullseye[0] - x, picture.bullseye[1] - y
    bearing = round(math.degrees(math.atan2(dy, dx))) % 360 or 360
    return t(
        "campaign.mission_deck.bullseye",
        dms=to_dms(lat, lon),
        bearing=f"{bearing:03d}",
        range=round(math.hypot(dx, dy) / NM),
        base=name,
    )


def bullseye_text(picture: MissionPicture, x: float, y: float) -> str:
    """``Bullseye : 050° / 18.3 nm``: where a point stands from the bullseye, on the grid the legs are given on.

    The hand-made briefing gave it on every target and every step (FIX-BRIEFING-PARITY-MARTEAU 04).
    """
    dx, dy = x - picture.bullseye[0], y - picture.bullseye[1]
    distance = math.hypot(dx, dy) / NM
    bearing = round(math.degrees(math.atan2(dy, dx))) % 360 if distance >= 0.05 else 0
    return t("mission_briefing.bullseye", bearing=f"{bearing:03d}", nm=f"{distance:.1f}")


#: How far from a zone's centre a waypoint still names it, at least, metres; the zone's radius when larger.
_LABEL_REACH = 1_500


def objective_waypoint(area: ObjectiveArea, picture: MissionPicture) -> str:
    """The flight plan's point on an objective (``KHAN``): the nearest within its zone; empty when none is.

    The maps and the pages name an objective as the pilots' route does, not by its trigger zone
    (``EWRTellHara`` for ``TELLHARA``, FIX-BRIEFING-PARITY-MARTEAU 02).
    """
    reach = max(area.radius, _LABEL_REACH)
    near = [
        (math.dist((wp.x, wp.y), (area.x, area.y)), wp.name)
        for plan in picture.flight_plans
        for wp in plan.waypoints
        if (wp.x, wp.y) != picture.bullseye
    ]
    inside = [found for found in near if found[0] <= reach]
    return min(inside)[1] if inside else ""


def _placed(
    areas: list[tuple[Objective, ObjectiveArea]], picture: MissionPicture, folder: Path
) -> list[tuple[Objective, ObjectiveArea]]:
    """Each objective's area with its label and, when the mission has an operation, whether it is secondary."""
    roles = operation_zones(folder)
    placed = []
    for objective, area in areas:
        # a chained zone (a counter-attack) is neither main nor secondary: it does not count in the operation's end
        secondary = roles is not None and area.zone.lower() not in roles[0] | roles[1]
        placed.append((objective, replace(area, label=objective_waypoint(area, picture), secondary=secondary)))
    return placed


def _objective_title(objective: Objective, area: ObjectiveArea) -> str:
    """``KHAN — La ligne de Khan Arnabah``, and ``(secondaire)`` after a secondary objective."""
    title = t("mission_briefing.title.objective", label=area.name, title=objective.title)
    return title + (t("mission_briefing.secondary_mark") if area.secondary else "")


def _threat_line(threat: Threat) -> str:
    """``SA-6 Gori — SA-6 : engagement ≈ 13 nm, détection ≈ 30 nm (source)``, with what is known."""
    name = f"{threat.name} — {threat.family}" if threat.family else threat.name
    ranges = []
    if threat.engagement_nm:
        ranges.append(t("mission_briefing.engagement", nm=f"{threat.engagement_nm:g}"))
    if threat.detection_nm:
        ranges.append(t("mission_briefing.detection", nm=f"{threat.detection_nm:g}"))
    if not ranges:
        return name
    source = f" ({threat.source})" if threat.source else ""
    return line(name, ", ".join(ranges) + source)


#: A waypoint number as the texts write it: ``W3``, ``WP3``, ``W 3``.
_WAYPOINT_REFERENCE = re.compile(r"\bW(?:P)?\s?(\d{1,2})\b")


def _folded(text: str) -> str:
    """Lower case, accents removed: ``Nœud`` and ``noeud`` compare equal."""
    return (
        unicodedata.normalize("NFKD", text.replace("œ", "oe").replace("Œ", "OE"))
        .encode("ascii", "ignore")
        .decode()
        .lower()
    )


def _waypoint_warnings(content: BriefingContent, picture: MissionPicture) -> list[str]:
    """Compare the ``W<n>`` of the assistant's words with the route the built mission carries.

    A route gains a point (the build's BULLSEYE, one more objective) after the texts were written, and
    every number after it is false: measured on a run whose briefing and F10 labels said W3, W4, W6 for
    points that were W4, W5, W7 (FIX-DEMO-VIDEO-FINDINGS 07). Nobody without PowerPoint can re-read the
    rendered pages, so the numbers are checked here: one past the end of the longest route, or one the
    same line contradicts by naming another point of the first route.

    Args:
        content: What the assistant wrote.
        picture: What the built mission says.

    Returns:
        One warning per number that does not fit; none when the mission has no route.
    """
    if not picture.flight_plans:
        return []
    longest = max(len(plan.waypoints) for plan in picture.flight_plans)
    names = [wp.name for wp in picture.flight_plans[0].waypoints]
    lines = [
        *content.context,
        *content.mission,
        *content.air_opposition,
        *content.air_defences,
        *content.other,
        *(text for objective in content.objectives for text in (objective.title, *objective.task)),
    ]
    warnings = []
    for text in lines:
        numbers = {int(found) for found in _WAYPOINT_REFERENCE.findall(text)}
        for number in sorted(numbers):
            if number > longest:
                warnings.append(
                    f"Waypoint W{number} in {text!r}: the route has {longest} point(s) after the departure "
                    "(counting the BULLSEYE the build adds); numbers past it point at nothing."
                )
        if len(numbers) == 1:
            number = next(iter(numbers))
            for index, name in enumerate(names, 1):
                if index != number and re.search(rf"\b{re.escape(_folded(name))}\b", _folded(text)):
                    warnings.append(
                        f"Waypoint {name} is W{index} on the route, not W{number} as {text!r} says: "
                        "numbers shift when a point (the BULLSEYE) is declared before the one named."
                    )
    return warnings


def _boxes(names: list[str], count: int) -> str:
    """``1. Zip  2. ______  3. ______  4. ______``: the pilots' names, a box for each seat not filled."""
    return "  ".join(f"{n}. {names[n - 1] if n <= len(names) else '______'}" for n in range(1, (count or 4) + 1))


def _package_rows(content: BriefingContent, picture: MissionPicture, warnings: list[str]) -> list[tuple[str, ...]]:
    """The ATO's flights, one row each: the mission's flight groups, then the flights on dynamic slots.

    A base lists the flights the assistant names, never the types it offers: 61 on a Syria base made six
    slides of ATO for two bases (FEAT-OBJECTIVE-BRIEFING-QUALITY 01). A base offering slots that no flight
    names gets one row of empty boxes.
    """
    free = t("campaign.mission_deck.free_loadout")
    rows: list[tuple[str, ...]] = []
    for group in picture.flights:
        start = group.base or t("campaign.mission_deck.flight_air_start")
        boxes = _boxes(content.pilots.get(group.name, []), group.count)
        aircraft = f"{group.count} × {group.aircraft}"
        rows.append((group.callsign, aircraft, start, "", boxes, content.loadouts.get(group.name, free)))
    bases = {airfield.name.lower(): airfield.name for airfield in picture.airfields}
    used: set[str] = set()
    for flight in content.flights:
        base = bases.get(flight.base.lower())
        if flight.base and base is None:
            warnings.append(
                f"Flight {flight.callsign!r}: {flight.base!r} offers no dynamic slot to the players' side; "
                "listed under any base."
            )
        used.add(base or "")
        aircraft = f"{flight.count} × {flight.aircraft}" if flight.count and flight.aircraft else flight.aircraft
        where = t("mission_briefing.slot_base", base=base) if base else t("mission_briefing.flights_any_base")
        boxes = _boxes(content.pilots.get(flight.callsign, []), flight.count)
        rows.append((flight.callsign, aircraft, where, flight.task, boxes, content.loadouts.get(flight.callsign, free)))
    for airfield in picture.airfields:
        if airfield.name not in used:
            rows.append((t("mission_briefing.ato.open_slots"), "", airfield.name, "", _boxes([], 4), free))
    return rows


def _ato_page(content: BriefingContent, picture: MissionPicture, warnings: list[str]) -> TablePage:
    """The ATO as tables: the package, the support, the control (FIX-BRIEFING-PARITY-MARTEAU 07)."""
    tables = []
    package = _package_rows(content, picture, warnings)
    if package:
        columns = tuple(
            t(f"mission_briefing.ato.{name}") for name in ("flight", "aircraft", "base", "task", "pilots", "loadout")
        )
        tables.append(
            Table(
                t("campaign.mission_deck.heading.package"),
                columns,
                (0.11, 0.15, 0.17, 0.17, 0.25, 0.15),
                tuple(package),
            )
        )
    support = [
        (
            aircraft.callsign,
            t(f"campaign.mission_deck.role.{aircraft.role}"),
            aircraft.aircraft,
            f"{mhz(aircraft.frequency)} MHz",
            aircraft.tacan or "—",
        )
        for aircraft in picture.support
    ]
    if support:
        columns = tuple(
            t(f"mission_briefing.ato.{name}") for name in ("callsign", "role", "type", "frequency", "tacan")
        )
        tables.append(
            Table(t("campaign.mission_deck.heading.support"), columns, (0.2, 0.2, 0.25, 0.2, 0.15), tuple(support))
        )
    control = [(a.name, f"{mhz(a.uhf)} MHz", f"{mhz(a.vhf)} MHz", a.tacan or "—") for a in picture.airfields]
    if control:
        columns = tuple(t(f"mission_briefing.ato.{name}") for name in ("airfield", "uhf", "vhf", "tacan"))
        tables.append(Table(t("campaign.mission_deck.heading.control"), columns, (0.4, 0.2, 0.2, 0.2), tuple(control)))
    carriers = tuple(carrier_line(carrier) for carrier in picture.carriers)
    after = [Block(None if control else t("campaign.mission_deck.heading.control"), carriers)] if carriers else []
    return TablePage(t("campaign.mission_deck.title.ato"), tables, after)


def _situation_page(
    content: BriefingContent,
    picture: MissionPicture,
    starts: list[tuple[str, float, float]],
) -> Page:
    blocks = [Block(t("campaign.mission_deck.heading.context"), content.context)] if content.context else []
    if content.mission:
        items = content.mission if content.mission_prose else tuple(f"– {item}" for item in content.mission)
        blocks.append(Block(t("campaign.mission_deck.heading.mission"), items))
    start = starts[0] if starts else None
    if start is None and picture.carriers:
        start = (picture.carriers[0].name, picture.carriers[0].x, picture.carriers[0].y)
    blocks.append(
        Block(t("campaign.mission_deck.heading.bullseye"), (_bullseye_line(picture.theatre, picture, start),))
    )
    departure = [
        t("campaign.mission_deck.departure_carrier", carrier=c.name, heading=f"{c.heading:03d}")
        for c in picture.carriers
    ]
    dynamic = {airfield.name for airfield in picture.airfields}
    for name, _, _ in starts:
        flights = [flight.callsign for flight in picture.flights if flight.base == name]
        if flights:
            departure.append(line(name, _join(flights)))
        elif name in dynamic:
            departure.append(t("campaign.mission_deck.departure_airfield", airfield=name))
    if departure:
        blocks.append(Block(t("campaign.mission_deck.heading.departure"), tuple(departure)))
    threats = [f"– {_threat_line(threat)}" for threat in content.threats]
    if picture.qra_zones:
        threats.append(f"– {t('mission_briefing.qra')}")
    if threats:
        blocks.append(Block(t("campaign.mission_deck.heading.threat"), tuple(threats)))
    blocks.append(Block(t("campaign.mission_deck.heading.weather"), weather_lines(picture)))
    date = t("campaign.mission_deck.date", day=picture.date.day, month=picture.date.month, year=picture.date.year)
    blocks.append(
        Block(
            t("campaign.mission_deck.heading.time"),
            (t("campaign.mission_deck.time", date=date, time=clock(picture.start_time)),),
        )
    )
    return Page(t("campaign.mission_deck.title.situation"), blocks)


def _prose_page(title: str, paragraphs: tuple[str, ...]) -> Page | None:
    """A page of the assistant's prose under its title; ``None`` when it wrote none."""
    return Page(title, [Block(None, paragraphs)]) if paragraphs else None


def _coordination_page(content: BriefingContent) -> TablePage | None:
    """The package's code words, each with who acts on it, as a table; ``None`` when there is none.

    FEAT-OBJECTIVE-MISSION-V2 03, a table since FIX-BRIEFING-PARITY-MARTEAU 07.
    """
    if not content.code_words:
        return None
    columns = tuple(t(f"mission_briefing.code_words.{name}") for name in ("word", "meaning", "who"))
    rows = tuple((code.word, code.meaning, code.who) for code in content.code_words)
    return TablePage(t("mission_briefing.title.coordination"), [Table(None, columns, (0.2, 0.42, 0.38), rows)])


def _deployment_page(content: BriefingContent) -> TablePage | None:
    """The deployment: its ``<when> — <who> : <what>`` paragraphs as a timeline table, the others after it.

    ``None`` when the assistant wrote none (FIX-BRIEFING-PARITY-MARTEAU 07).
    """
    if not content.deployment:
        return None
    rows, prose = timeline(content.deployment)
    columns = tuple(t(f"mission_briefing.deployment.{name}") for name in ("when", "who", "what"))
    tables = [Table(None, columns, (0.17, 0.18, 0.65), tuple(rows))] if rows else []
    after = [Block(None, tuple(prose))] if prose else []
    return TablePage(t("mission_briefing.title.deployment"), tables, after)


def _code_word_blocks(content: BriefingContent) -> list[tuple[str | None, list[str]]]:
    """The code words as the kneeboard writes them: the word as a heading, its meaning and who acts under it."""
    return [(code.word, [code.meaning, *([f"→ {code.who}"] if code.who else [])]) for code in content.code_words]


def _attack_origin(picture: MissionPicture, starts: list[tuple[str, float, float]]) -> tuple[float, float] | None:
    """Where the players come from: their departure, their carrier, or the first point of their route."""
    if starts:
        return starts[0][1], starts[0][2]
    if picture.carriers:
        return picture.carriers[0].x, picture.carriers[0].y
    if picture.flight_plans and picture.flight_plans[0].waypoints:
        first = picture.flight_plans[0].waypoints[0]
        return first.x, first.y
    return None


def _pairs(theatre: str, photos: list[ReconPhoto]) -> list[list[tuple[Path, str]]]:
    """The photos two by two, each with its caption: what a slide or a kneeboard page holds."""
    captioned = [(photo.image, caption(theatre, photo)) for photo in photos]
    return [captioned[i : i + 2] for i in range(0, len(captioned), 2)]


def _objective_blocks(areas: list[tuple[Objective, ObjectiveArea]], chained: set[str]) -> list[Block]:
    """The objectives, the operation's first and the others after; a chained zone in neither list.

    Without an operation every objective is listed, as before (FIX-BRIEFING-PARITY-MARTEAU 05).
    """
    kept = [(objective, area) for objective, area in areas if area.zone.lower() not in chained]
    main = [objective.title for objective, area in kept if not area.secondary]
    other = [objective.title for objective, area in kept if area.secondary]
    if not other:
        heading = t("campaign.mission_deck.heading.objectives")
        return [Block(heading, tuple(f"{n}. {title}" for n, title in enumerate(main, 1)))] if main else []
    blocks = []
    if main:
        blocks.append(
            Block(
                t("mission_briefing.heading.main_objectives"),
                tuple(f"{n}. {title}" for n, title in enumerate(main, 1)),
            )
        )
    blocks.append(Block(t("mission_briefing.heading.secondary_objectives"), tuple(f"– {title}" for title in other)))
    return blocks


def _flow_page(
    content: BriefingContent,
    picture: MissionPicture,
    starts: list[tuple[str, float, float]],
    areas: list[tuple[Objective, ObjectiveArea]] | None = None,
    chained: set[str] | None = None,
) -> Page:
    blocks = []
    if areas is not None:
        blocks += _objective_blocks(areas, chained or set())
    elif content.objectives:
        blocks.append(
            Block(
                t("campaign.mission_deck.heading.objectives"),
                tuple(f"{n}. {objective.title}" for n, objective in enumerate(content.objectives, 1)),
            )
        )
    air = list(content.air_opposition) or ([t("mission_briefing.qra")] if picture.qra_zones else [])
    blocks.append(Block(t("campaign.mission_deck.heading.air"), tuple(air) or (t("campaign.mission_deck.no_air"),)))
    defences = [_threat_line(threat) for threat in content.threats] + list(content.air_defences)
    if defences:
        blocks.append(Block(t("campaign.mission_deck.heading.air_defence"), tuple(defences)))
    others = list(content.other)
    tankers = [s.callsign for s in picture.support if s.role != "awacs"]
    if tankers:
        others.append(t("campaign.mission_deck.refuelling", tankers=_join(tankers)))
    # the departure is no diversion: the deck said "Diversion fields: Senaki" for a flight leaving
    # Senaki (FIX-COWORK-SESSION-FINDINGS 04); the assistant may name them instead
    departed = {name for name, _, _ in starts}
    diversions = (
        list(content.diversions)
        if content.diversions is not None
        else [a.name for a in picture.airfields if a.name not in departed]
        + [c.name for c in picture.carriers if c.name not in departed]
    )
    if diversions:
        others.append(t("campaign.mission_deck.diversion", fields=_join(diversions)))
    if picture.plane_guard:
        others.append(t("campaign.mission_deck.plane_guard"))
    if picture.csar:
        others.append(t("campaign.mission_deck.csar"))
    if others:
        blocks.append(Block(t("campaign.mission_deck.heading.other"), tuple(others)))
    return Page(t("campaign.mission_deck.title.flow"), blocks)


def _normalised(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.lower())


def flies(aircraft: str, unit_type: str) -> bool:
    """Whether a flight's aircraft, as the assistant writes it (``F/A-18C``, ``AH-64D, OH-58D``), names a DCS type.

    Written without the type's suffix: ``F-16C`` is ``F-16C_50``, ``A-10C`` both ``A-10C`` and ``A-10C_2``.
    """
    kind = _normalised(unit_type)
    for part in re.split(r",|\bou\b|\bor\b|\+", aircraft):
        name = _normalised(part)
        if len(name) >= 3 and (kind.startswith(name) or name.startswith(kind)):
            return True
    return False


def _flown(flights: tuple[PackageFlight, ...], types: list[str]) -> dict[str, list[str]]:
    """The DCS types the ATO's flights fly, each with the callsigns flying it."""
    found: dict[str, list[str]] = {}
    for unit_type in types:
        callsigns = [flight.callsign for flight in flights if flies(flight.aircraft, unit_type)]
        if callsigns:
            found[unit_type] = callsigns
    return found


def _ato_plans(
    plans: list[RadioPlan], flights: tuple[PackageFlight, ...], groups: frozenset[str] = frozenset()
) -> list[RadioPlan] | None:
    """The radio plans of the types the ATO flies, one per distinct channel list, named by their flights.

    A base offering dynamic slots carries the presets of every type it offers: 61 on a Syria base made a
    frequency plan of seven pages, each type's radios listed (FIX-BRIEFING-PARITY-MARTEAU 06). A plan
    with no named channel (an FM radio's thirty bare frequencies) says nothing a pilot needs.

    A flight group the mission places (``groups``, by name) keeps its plans under its own name: a mission
    may mix named slots and dynamic ones (pr-code-review of #1210).

    Returns:
        The plans, ``None`` when the ATO names no flight or none of its aircraft is a type of the plans.
    """
    if not flights:
        return None
    flown = _flown(
        flights, [name for name in dict.fromkeys(n for plan in plans for n in plan.flights) if name not in groups]
    )
    if not flown:
        return None
    flown.update({name: [name] for plan in plans for name in plan.flights if name in groups})
    merged: dict[tuple[tuple[int, str, float], ...], tuple[list[int], list[str]]] = {}
    for plan in plans:
        carriers = [name for name in plan.flights if name in flown]
        if not carriers or not any(name for _, name, _ in plan.channels):
            continue
        radios, callsigns = merged.setdefault(plan.channels, ([], []))
        if plan.radio not in radios:
            radios.append(plan.radio)
        for name in carriers:
            callsigns += [callsign for callsign in flown[name] if callsign not in callsigns]
    order = [flight.callsign for flight in flights] + sorted(groups)
    return [
        RadioPlan(radios[0], tuple(sorted(callsigns, key=order.index)), channels, tuple(radios))
        for channels, (radios, callsigns) in merged.items()
    ]


def _frequencies_page(picture: MissionPicture, plans: list[RadioPlan], flights: tuple[PackageFlight, ...] = ()) -> Page:
    """The presets the build put in the cockpits, radio by radio; the mission's channels without them.

    With the ATO's flights, only the plans their aircraft carry, one block each (:func:`_ato_plans`).
    """
    if not plans:
        return frequencies_page(picture)
    plans = _ato_plans(plans, flights, frozenset(flight.name for flight in picture.flights)) or plans
    blocks = []
    for plan in plans:
        lines = [
            t("mission_briefing.channel", channel=channel, name=name, frequency=mhz(freq))
            for channel, name, freq in plan.channels
            if name
        ]
        # a channel with no name is a bare frequency (an FM radio's thirty): one line for all of them
        unnamed = [
            t("mission_briefing.channel_unnamed", channel=channel, frequency=mhz(freq))
            for channel, name, freq in plan.channels
            if not name
        ]
        if unnamed:
            lines.append(" · ".join(unnamed))
        radio = " / ".join(str(number) for number in plan.radios) if plan.radios else str(plan.radio)
        blocks.append(Block(t("mission_briefing.radio", radio=radio, flights=", ".join(plan.flights)), tuple(lines)))
    control = [carrier_line(carrier) for carrier in picture.carriers]
    if control:
        blocks.append(Block(t("campaign.mission_deck.heading.control"), tuple(control)))
    return Page(t("campaign.mission_deck.title.frequencies"), blocks)


def _feet(metres: float) -> str:
    """``2 640`` feet, the thousands set apart as the pages write them."""
    return f"{metres_to_feet(metres):,.0f}".replace(",", " ")


def _target_sheets(
    picture: MissionPicture,
    areas: list[tuple[Objective, ObjectiveArea]],
    zooms: list[TacticalMapReport],
    warnings: list[str],
) -> CardSection:
    """One sheet per target: what it is and does, its coordinates in DMS and DDM, its elevation, its bullseye,
    a thumbnail of its zoom.

    The list of names and DMS it replaces was arid (FEAT-OBJECTIVE-BRIEFING-QUALITY 03); a target is a unit
    when its units carry names of their own (FIX-BRIEFING-PARITY-MARTEAU 03).
    """
    theatre = picture.theatre
    grid = grid_for_theatre(theatre)
    if grid is None and any(area.targets for _, area in areas):
        warnings.append(
            f"No terrain elevation grid for {theatre} on this PC: the target coordinates carry no altitude "
            "(see the terrain_elevation action)."
        )
    section = CardSection(t("mission_briefing.title.targets"))
    for (objective, area), zoom in zip(areas, zooms, strict=True):
        cards = []
        for target in area.targets:
            lat, lon = xy_to_latlon(theatre, target.x, target.y)
            lines = [target.kind] if target.kind else []
            lines += [target.role] if target.role else []
            lines += [f"DMS {to_dms(lat, lon)}", f"DDM {to_ddm(lat, lon)}"]
            height = grid.elevation_at(target.x, target.y) if grid else None
            if height is not None:
                lines.append(t("mission_briefing.elevation", feet=_feet(height), metres=round(height)))
            lines.append(bullseye_text(picture, target.x, target.y))
            thumb = thumbnail(zoom.path, zoom.locate(target.x, target.y)) if zoom.locate else None
            cards.append(Card(target.name, tuple(lines), thumb))
        heading = _objective_title(objective, area)
        section.groups.append((heading, cards or [Card(t("mission_briefing.no_target"), ())]))
    return section


def _step_sheets(
    content: BriefingContent,
    picture: MissionPicture,
    starts: list[tuple[str, float, float]],
    overview: TacticalMapReport,
    warnings: list[str],
) -> CardSection | None:
    """One sheet per waypoint, as a DCS kneeboard shows a route: the leg that reaches it, where it is, what is done there.

    Returns:
        The section, ``None`` when the players carry no route.
    """
    theatre = picture.theatre
    notes = {name.lower(): text for name, text in content.waypoint_notes.items()}
    by_plan = {
        category: {name.lower(): text for name, text in plan_notes.items()}
        for category, plan_notes in content.plan_notes.items()
    }
    named = {wp.name.lower() for plan in picture.flight_plans for wp in plan.waypoints}
    for name in content.waypoint_notes:
        if name.lower() not in named:
            warnings.append(
                f"waypoint_notes: no waypoint named {name!r} on the players' routes; the note is not shown."
            )
    for category, plan_notes in content.plan_notes.items():
        on_plan = {
            wp.name.lower() for plan in picture.flight_plans if plan.category == category for wp in plan.waypoints
        }
        for name in plan_notes:
            if name.lower() not in on_plan:
                warnings.append(
                    f"waypoint_notes.{category}: no waypoint named {name!r} on the {category}s' route; "
                    "the note is not shown."
                )
    section = CardSection(t("mission_briefing.title.steps"))
    for plan in picture.flight_plans:
        heading = t(f"campaign.mission_deck.heading.{plan.category}s")
        if sum(1 for other in picture.flight_plans if other.category == plan.category) > 1:
            heading = f"{heading} — {', '.join(plan.aircraft)}"
        cards = []
        for step in route_steps(theatre, plan, starts):
            if step.waypoint is None:
                continue
            lat, lon = xy_to_latlon(theatre, step.x, step.y)
            lines = [t("mission_briefing.step_from", origin=step.origin, leg=leg_text(step))] if step.origin else []
            lines += [f"DMS {to_dms(lat, lon)}", f"DDM {to_ddm(lat, lon)}"]
            lines.append(t("mission_briefing.step_altitude", altitude=altitude(step.waypoint)))
            lines.append(bullseye_text(picture, step.x, step.y))
            note = by_plan.get(plan.category, {}).get(step.name.lower()) or notes.get(step.name.lower())
            if note:
                lines.append(note)
            thumb = thumbnail(overview.path, overview.locate(step.x, step.y)) if overview.locate else None
            cards.append(Card(f"W{step.number} — {step.name}", tuple(lines), thumb))
        section.groups.append((heading, cards))
    return section if section.groups else None


def objective_notes(objective: Objective, area: ObjectiveArea, picture: MissionPicture) -> list[tuple[str, bool]]:
    """What an objective's page says beside its zoom: its task, for whom, its bullseye, its point, its targets.

    The hand-made briefing gave "Pour", "Bullseye" and "Point" on every objective (FIX-BRIEFING-PARITY-MARTEAU 04).
    Each note is ``(text, is_heading)``.
    """
    notes = [(objective.title, True), *[(text, False) for text in objective.task]]
    if objective.flights:
        notes.append((t("mission_briefing.for", flights=_join(list(objective.flights))), False))
    notes.append((bullseye_text(picture, area.x, area.y), False))
    if area.label:
        notes.append((t("mission_briefing.waypoint", name=area.label), False))
    if area.targets:
        notes += [(t("mission_briefing.heading.targets"), True), (_join([x.name for x in area.targets]), False)]
    return notes


def _cover_date_line(picture: MissionPicture) -> str:
    """The cover's line under the subtitle: ``13 octobre 2015 — 06h15 locale``."""
    date = t("campaign.mission_deck.date", day=picture.date.day, month=picture.date.month, year=picture.date.year)
    return f"{date} — {t('campaign.mission_deck.local', time=clock(picture.start_time))}"


def _write_cover(deck: PresentationType, title: str, picture: MissionPicture, contents: list[str]) -> None:
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = _frame(slide, _MARGIN, Emu(1700000), Emu(6800000), Emu(3200000))
    _run(frame.paragraphs[0], title, 54, bold=True, colour=_BLUE)
    _run(frame.add_paragraph(), t("campaign.mission_deck.subtitle"), 22, colour=_GREY)
    _run(frame.add_paragraph(), _cover_date_line(picture), 16, italic=True, colour=_GREY)
    toc = _frame(slide, Emu(7900000), Emu(1500000), Emu(3900000), Emu(4600000))
    _run(toc.paragraphs[0], t("campaign.deck.contents"), 20, bold=True)
    for item in contents:
        _run(toc.add_paragraph(), item, 14)


def _wide_picture_page(deck: PresentationType, title: str, image: Path) -> None:
    """A map under its title, as large as the slide's body holds it, centred."""
    slide = _titled_slide(deck, title)
    with Image.open(image) as opened:
        ratio = opened.width / opened.height
    top = Emu(1000000)
    room_w, room_h = _WIDTH - 2 * _MARGIN, _HEIGHT - top - Emu(200000)
    width = min(room_w, int(room_h * ratio))
    height = int(width / ratio)
    left = _MARGIN + (room_w - width) // 2
    slide.shapes.add_picture(str(image), Emu(left), top, width=Emu(width), height=Emu(height))


def _photo_slide(deck: PresentationType, title: str, photos: list[tuple[Path, str]]) -> None:
    """One or two recon photos side by side under the title, each with its caption under it."""
    slide = _titled_slide(deck, title)
    gap = Emu(200000)
    top = Emu(1000000)
    caption_h = Emu(700000)
    room_w = (_WIDTH - 2 * _MARGIN - (len(photos) - 1) * gap) // len(photos)
    room_h = _HEIGHT - top - caption_h - Emu(200000)
    for index, (image, text) in enumerate(photos):
        with Image.open(image) as opened:
            ratio = opened.width / opened.height
        width = min(room_w, int(room_h * ratio))
        height = int(width / ratio)
        left = _MARGIN + index * (room_w + gap) + (room_w - width) // 2
        slide.shapes.add_picture(str(image), Emu(left), top, width=Emu(width), height=Emu(height))
        _run(_frame(slide, left, top + height + Emu(50000), width, caption_h).paragraphs[0], text, 12, colour=_GREY)


def _jpeg(png: Path) -> bytes:
    """A map as a DCS briefing picture: JPEG, brought down to :data:`_PICTURE_WIDTH` when wider."""
    with Image.open(png) as image:
        picture = image.convert("RGB")
    if picture.width > _PICTURE_WIDTH:
        picture = picture.resize((_PICTURE_WIDTH, round(picture.height * _PICTURE_WIDTH / picture.width)))
    buffer = io.BytesIO()
    picture.save(buffer, format="JPEG", quality=90)
    return buffer.getvalue()


def _list_pictures(mission: DcsMission, pictures: dict[str, str], side: str) -> list[str]:
    """Replace this tool's pictures in the players' side's and the neutral lists; return the other side's list."""
    content = mission.mission_content
    if content is None:
        raise ValueError("Not a valid DCS mission (missing 'mission' content)")
    resources = {k: v for k, v in (mission.map_resource_content or {}).items() if not str(k).startswith(PICTURE_PREFIX)}
    resources.update(pictures)
    mission.map_resource_content = resources
    for key in (_PICTURE_KEYS[side], _PICTURE_KEYS["neutral"]):
        kept = [str(entry) for entry in indexed(content.get(key)) if not str(entry).startswith(PICTURE_PREFIX)]
        content[key] = [*pictures, *kept]
    other = _PICTURE_KEYS["red" if side == "blue" else "blue"]
    return [str(entry) for entry in indexed(content.get(other))]


#: A picture file this tool writes: what it may delete on a later run, and nothing else.
_GENERATED_PICTURE = re.compile(r"briefing-[A-Za-z0-9._-]+\.jpg")


def _stale_pictures(l10n: Path, previous: set[str], kept: set[str]) -> list[Path]:
    """Return the pictures an earlier run wrote and this one no longer writes, safe to delete.

    The names come from the mission's `mapResource`, and a mission may come from anywhere: only a
    name this tool generates, resolved inside `l10n/DEFAULT/`, is returned (security review, 6.31.0).

    Args:
        l10n: The mission folder's `l10n/DEFAULT/`.
        previous: The file names the `MCP_Briefing_` resources pointed at.
        kept: The file names this run writes.

    Returns:
        The files to delete, sorted.
    """
    root = l10n.resolve()
    stale = []
    for name in sorted(previous - kept):
        if _GENERATED_PICTURE.fullmatch(name) is None:
            continue
        path = (root / name).resolve()
        if path.parent == root:
            stale.append(path)
    return stale


def built_archives(folder: Path, miz: Path) -> list[Path]:
    """The build's own `.miz`, then each weather variant the build wrote into `missions/`.

    The players fly a variant, not the build's own archive: pictures written into the latter only
    reached them after a rebuild (FIX-COWORK-SESSION-FINDINGS 01).

    Args:
        folder: The mission folder.
        miz: The build's own `.miz` (`find_built_mission`).

    Returns:
        Every archive the briefing goes into, the build's own first.
    """
    return [miz, *sorted((folder / VARIANTS_FOLDER).glob("*.miz"))]


#: The pictures a kneeboard folder holds: what replacing it moves away, and nothing else.
_KNEEBOARD_PICTURE = re.compile(r".+\.(?:jpe?g|png)", re.IGNORECASE)

#: Where replaced kneeboard pages are kept, in the mission folder.
_KNEEBOARD_BACKUP = Path(".veaf-backups")


#: The build's own kneeboard page, the presets' calibration: not a hand-made page, never taken away.
_CALIBRATION = KNEEBOARD_CALIBRATION_PATH.rsplit("/", 1)[-1]


def _foreign_page(name: str, pages: dict[str, bytes]) -> bool:
    """A picture of the common kneeboard neither this tool nor the build wrote: what ``kneeboard_replace`` takes away."""
    if name == _CALIBRATION or name in pages or is_kneeboard_page(name):
        return False
    return _KNEEBOARD_PICTURE.fullmatch(name) is not None


def _named_plan(plan: FlightPlan, flights: tuple[PackageFlight, ...]) -> FlightPlan:
    """A dynamic-slot plan named by the ATO's flights that fly it, not by the 61 types that carry it."""
    callsigns = [flight.callsign for flight in flights if any(flies(flight.aircraft, kind) for kind in plan.aircraft)]
    return replace(plan, aircraft=tuple(callsigns)) if callsigns else plan


def _write_kneeboard_folder(folder: Path, pages: dict[str, bytes], replace_others: bool = False) -> list[str]:
    """Put the kneeboard pages in the mission folder, so a rebuild carries them; drop the ones no longer written.

    With ``replace_others``, the pictures of the common kneeboard this tool did not write are moved into
    ``.veaf-backups/kneeboard-<date>/``: a hand-made kneeboard beside the tool's gave the players both
    (FIX-BRIEFING-PARITY-MARTEAU 09).

    Returns:
        The names of the pages moved away.
    """
    target = _folder_l10n(folder).parent.parent / KNEEBOARD_FOLDER
    moved: list[str] = []
    if target.is_dir():
        for stale in sorted(target.iterdir()):
            if is_kneeboard_page(stale.name) and stale.name not in pages:
                stale.unlink(missing_ok=True)
            elif replace_others and stale.is_file() and _foreign_page(stale.name, pages):
                backup = folder / _KNEEBOARD_BACKUP / f"kneeboard-{datetime.now():%Y%m%d-%H%M%S}"
                backup.mkdir(parents=True, exist_ok=True)
                shutil.move(str(stale), str(backup / stale.name))
                moved.append(stale.name)
    if pages:
        target.mkdir(parents=True, exist_ok=True)
    for name, data in pages.items():
        (target / name).write_bytes(data)
    return moved


def _write_pictures(
    folder: Path,
    archives: list[Path],
    maps: list[Path],
    side: str,
    kneeboard: dict[str, bytes],
    replace_others: bool = False,
) -> tuple[list[str], list[str], list[str]]:
    """Put the maps and the kneeboard pages in the mission folder and every built `.miz`, backed up first.

    Returns:
        The resource keys, in display order, the other side's pictures found, and the kneeboard pages that
        were not the tool's and were taken away (``replace_others``).
    """
    files = {f"briefing-{path.stem}.jpg": _jpeg(path) for path in maps}
    pictures = {f"{PICTURE_PREFIX}{path.stem}": f"briefing-{path.stem}.jpg" for path in maps}

    mission = load_folder_mission(folder)
    previous = {str(v) for k, v in (mission.map_resource_content or {}).items() if str(k).startswith(PICTURE_PREFIX)}
    other = _list_pictures(mission, pictures, side)
    l10n = _folder_l10n(folder)
    l10n.mkdir(parents=True, exist_ok=True)
    for stale in _stale_pictures(l10n, previous, set(files)):
        stale.unlink(missing_ok=True)
    for name, data in files.items():
        (l10n / name).write_bytes(data)
    save_folder_mission(mission, folder)
    moved = _write_kneeboard_folder(folder, kneeboard, replace_others)

    for miz in archives:
        built = read_miz(miz)
        _list_pictures(built, pictures, side)
        additional: dict[str, Any] = {f"{DEFAULT_SCRIPTS_LOCATION}/{name}": data for name, data in files.items()}
        additional.update({f"{KNEEBOARD_FOLDER}/{name}": data for name, data in kneeboard.items()})
        if f"{DEFAULT_SCRIPTS_LOCATION}/mapResource" not in list_members(miz):
            # `write_miz` rewrites a `mapResource` member the archive has; one it lacks has to be supplied
            additional[f"{DEFAULT_SCRIPTS_LOCATION}/mapResource"] = (
                "mapResource = \n"
                + luadata.serialize(
                    built.map_resource_content, indent="  ", indent_level=0, always_provide_keyname=True, sort=True
                )
            ).encode()
        # a page an earlier run wrote and this one does not (one objective less, kneeboard off) leaves
        # the archive too, not only the folder
        stale_pages = set()
        for member in list_members(miz):
            if not member.startswith(f"{KNEEBOARD_FOLDER}/") or "/" in member[len(KNEEBOARD_FOLDER) + 1 :]:
                continue
            name = member.rsplit("/", 1)[-1]
            if is_kneeboard_page(name) and name not in kneeboard:
                stale_pages.add(member)
            elif replace_others and _foreign_page(name, kneeboard):
                stale_pages.add(member)
                if name not in moved:
                    moved.append(name)
        backup_before_write(miz)
        write_miz(built, miz, additional_files=additional, removed_files=stale_pages)
    return list(pictures), other, moved


def _inset(overview: TacticalMapReport, area: ObjectiveArea) -> Inset | None:
    """The overview as a zoom's inset, the zoom's place marked; ``None`` when the overview cannot place it."""
    if overview.locate is None or not overview.metres_per_pixel:
        return None
    radius = zoom_half(area) / overview.metres_per_pixel
    return Inset(overview.path, overview.locate(area.x, area.y), radius)


def _kneeboard_pages(
    folder: Path,
    picture: MissionPicture,
    areas: list[tuple[Objective, ObjectiveArea]],
    threats: list[tuple[Threat, tuple[float, float]]],
    starts: list[tuple[str, float, float]],
    tiles: dict[str, Any],
    content: BriefingContent,
    recon: dict[str, list[ReconPhoto]],
    ground: GroundPicture | None = None,
) -> dict[str, bytes]:
    """The briefing for the knee: the maps drawn for the 2:3 page, the flight card, the targets.

    The maps are turned so that the attack axis, from the players' departure to the objectives, points
    up, as the maps of "Marteau du Golan" were by hand; a north arrow says where north is. The combat
    area follows the overview, each objective's recon photos its zoom, and the code words close the set
    (FEAT-OBJECTIVE-MISSION-V2 02 to 04).

    Returns:
        Page file name -> JPEG, in the order the kneeboard shows them.
    """
    theatre = picture.theatre
    maps_folder = folder / MAPS_FOLDER / "kneeboard"
    origin = _attack_origin(picture, starts)
    turn = attack_axis_rotation(
        xy_to_latlon(theatre, *origin) if origin else None,
        [xy_to_latlon(theatre, area.x, area.y) for _, area in areas],
    )
    overview = render_overview_map(
        picture,
        areas,
        threats,
        starts,
        maps_folder / OVERVIEW_FILE,
        aspect=MAP_ASPECT,
        width=_KNEEBOARD_MAP,
        rotation=turn,
        ground=ground,
        **tiles,
    )
    sheets: list[tuple[str, bytes]] = [("overview", map_page(t("campaign.mission_deck.title.tactical"), overview.path))]
    if areas or threats:
        combat = render_overview_map(
            picture,
            areas,
            threats,
            starts,
            maps_folder / COMBAT_AREA_FILE,
            aspect=MAP_ASPECT,
            width=_KNEEBOARD_MAP,
            combat_area=True,
            rotation=turn,
            ground=ground,
            **tiles,
        )
        sheets.append(("combat-area", map_page(t("mission_briefing.title.combat_area"), combat.path)))
    for objective, area in areas:
        zoom = render_objective_zoom(
            theatre,
            objective,
            area,
            threats,
            maps_folder / zoom_file(area.zone),
            aspect=MAP_ASPECT,
            width=_KNEEBOARD_MAP,
            inset=_inset(overview, area),
            rotation=turn or 0.0,
            **tiles,
        )
        title = _objective_title(objective, area)
        stem = Path(zoom_file(area.zone)).stem
        sheets.append((stem, map_page(title, zoom.path)))
        recon_title = t("mission_briefing.title.recon", zone=area.name)
        for n, pair in enumerate(_pairs(theatre, recon.get(area.zone, [])), 1):
            sheets.append((f"{stem}-recon-{n}", photo_page(recon_title, pair)))
    card_blocks = flight_card_blocks(theatre, list(picture.flight_plans), starts)
    card = text_pages(t("mission_briefing.kneeboard.flight_card"), card_blocks) if card_blocks else []
    target_blocks = targets_blocks(
        theatre, [(_objective_title(objective, area), area) for objective, area in areas], picture.bullseye
    )
    targets = text_pages(t("mission_briefing.title.targets"), target_blocks) if target_blocks else []
    blocks = _code_word_blocks(content)
    codes = text_pages(t("mission_briefing.title.coordination"), blocks) if blocks else []
    for slug, pages in (("flight-card", card), ("targets", targets), ("code-words", codes)):
        sheets += [(slug if n == 1 else f"{slug}-{n}", page) for n, page in enumerate(pages, 1)]
    return {page_name(n, slug): data for n, (slug, data) in enumerate(sheets, 1)}


def build_mission_briefing(
    folder: Path,
    content: BriefingContent,
    *,
    side: str = "blue",
    source: TileSource = OSM,
    cache_dir: Path | None = None,
    fetch: Fetch | None = None,
    kneeboard: bool = True,
    recon_folder: Path | None = None,
    kneeboard_replace: bool = False,
) -> BriefingReport:
    """Write a one-shot mission's briefing deck and maps, and put the maps in its DCS briefing.

    Args:
        folder: The mission folder, built (`build_mission`).
        content: What the assistant writes: the scenario's words.
        side: The players' side.
        source: Where the map tiles come from.
        cache_dir: Where map tiles are kept between runs.
        fetch: How a map tile is downloaded (tests pass their own).
        kneeboard: Whether the maps, a flight card and the targets also go into every aircraft's kneeboard.
        recon_folder: The objectives' recon photos and their sidecars (`briefing_recon`), relative to the
            folder or absolute; the folder's `docs/recon/` when omitted, nothing when that does not exist.
        kneeboard_replace: Take the common kneeboard's pages this tool did not write out of the folder (kept
            in `.veaf-backups/`) and out of every built `.miz`: the tool's kneeboard replaces a hand-made one.

    Returns:
        The deck, its page count, the maps, the pictures' keys, the `.miz` read, and the warnings.

    Raises:
        ValueError: If the mission is not built, or an objective names no trigger zone of the mission.
    """
    miz = find_built_mission(folder)
    if miz is None:
        raise ValueError(f"No built mission in {folder}: run build_mission first, the briefing reads the built .miz.")
    with language(_mission_language(folder)):
        recon = folder / (recon_folder if recon_folder is not None else RECON_FOLDER)
        return _build(folder, miz, content, side, source, cache_dir, fetch, kneeboard, recon, kneeboard_replace)


def _slide_jpeg(png: Path, scratch: Path) -> Path:
    """A map as the deck and the PDF carry it: a JPEG, a tenth of the PNG's weight at that resolution.

    The overview's PNG at 2000 pixels and more weighs some 12 MB: embedded as it is, the deck reached 24 MB.
    """
    jpeg = scratch / f"{png.stem}.jpg"
    if not jpeg.is_file():
        with Image.open(png) as image:
            image.convert("RGB").save(jpeg, format="JPEG", quality=88)
    return jpeg


Sheet = Page | MapSheet | CardPage | PhotoSheet | TablePage


def _write_item(deck: PresentationType, item: Page | CardSection | TablePage, sheets: list[Sheet]) -> None:
    """A text page, a section of cards or a page of tables, on as many slides as it needs."""
    if isinstance(item, CardSection):
        for card_page in item.pages():
            write_card_page(deck, card_page)
            sheets.append(card_page)
    elif isinstance(item, TablePage):
        for n, table_page in enumerate(paginate_tables(item)):
            if n:
                table_page.title = t("campaign.deck.continued", title=item.title)
            write_table_page(deck, table_page)
            sheets.append(table_page)
    else:
        for part in paginate(item):
            _write_page(deck, part)
            sheets.append(part)


def _write_deck_and_pdf(
    folder: Path,
    content: BriefingContent,
    picture: MissionPicture,
    numbered: list[str],
    opening: list[Page | TablePage],
    closing: list[Page | CardSection | TablePage],
    maps_at: int,
    areas: list[tuple[Objective, ObjectiveArea]],
    overview: TacticalMapReport,
    zooms: list[TacticalMapReport],
    scratch: Path,
    combat: TacticalMapReport | None = None,
    recon: dict[str, list[ReconPhoto]] | None = None,
) -> tuple[int, int]:
    """Write the deck and the PDF of the same pages; return their page counts."""
    deck = Presentation()
    deck.slide_width, deck.slide_height = _WIDTH, _HEIGHT
    _write_cover(deck, content.title, picture, numbered)
    sheets: list[Sheet] = []
    for page in opening:
        _write_item(deck, page, sheets)
    # the key is drawn into the map: the slide gives the map its whole width
    _wide_picture_page(deck, numbered[maps_at], _slide_jpeg(overview.path, scratch))
    sheets.append(MapSheet(numbered[maps_at], _slide_jpeg(overview.path, scratch), []))
    first_zoom = maps_at + 1
    if combat is not None:
        _wide_picture_page(deck, numbered[first_zoom], _slide_jpeg(combat.path, scratch))
        sheets.append(MapSheet(numbered[first_zoom], _slide_jpeg(combat.path, scratch), []))
        first_zoom += 1
    zoom_titles = numbered[first_zoom : first_zoom + len(areas)]
    for title, (objective, area), zoom in zip(zoom_titles, areas, zooms, strict=True):
        notes = objective_notes(objective, area, picture)
        picture_page(deck, title, _slide_jpeg(zoom.path, scratch), notes)
        sheets.append(MapSheet(title, _slide_jpeg(zoom.path, scratch), notes))
        # its recon photos after the zoom, two to a slide, unnumbered (FEAT-OBJECTIVE-MISSION-V2 04)
        recon_title = t("mission_briefing.title.recon", zone=area.name)
        for pair in _pairs(picture.theatre, (recon or {}).get(area.zone, [])):
            _photo_slide(deck, recon_title, pair)
            sheets.append(PhotoSheet(recon_title, pair))
    for item in closing:
        _write_item(deck, item, sheets)
    out = folder / DECK_FILE
    out.parent.mkdir(parents=True, exist_ok=True)
    deck.save(str(out))
    pdf = folder / PDF_FILE
    cover = Cover(
        content.title,
        t("campaign.mission_deck.subtitle"),
        _cover_date_line(picture),
        t("campaign.deck.contents"),
        numbered,
    )
    return len(deck.slides), write_pdf(pdf, cover, sheets)


def _build(
    folder: Path,
    miz: Path,
    content: BriefingContent,
    side: str,
    source: TileSource,
    cache_dir: Path | None,
    fetch: Fetch | None,
    kneeboard: bool,
    recon_folder: Path,
    kneeboard_replace: bool = False,
) -> BriefingReport:
    built = read_miz(miz)
    table = built.mission_content or {}
    picture = read_mission_picture(miz, side, folder)
    dynamic = bool(picture.airfields)
    if not dynamic:  # the scaffold's dynamic-slot templates are nobody's cockpit here
        picture = replace(picture, flight_plans=_flight_plans(table, side, templates=False))
    elif content.flights:  # a dynamic plan is carried by every type a base offers: name it by the ATO's flights
        picture = replace(picture, flight_plans=[_named_plan(plan, content.flights) for plan in picture.flight_plans])
    theatre = picture.theatre
    enemy = "red" if side == "blue" else "blue"
    warnings: list[str] = _waypoint_warnings(content, picture)
    areas = _placed(
        [(objective, objective_area(table, objective.zone, enemy)) for objective in content.objectives], picture, folder
    )
    roles = operation_zones(folder)
    threats = []
    for threat in content.threats:
        position = group_position(table, threat.group)
        if position is None:
            warnings.append(f"Threat {threat.name!r}: no group named {threat.group!r} in the built mission, not drawn.")
        else:
            threats.append((threat, position))
    starts = departures(table, side, theatre, templates=dynamic)
    ground = ground_picture(table, side)

    maps_folder = folder / MAPS_FOLDER
    tiles: dict[str, Any] = {"cache_dir": cache_dir, "fetch": fetch, "source": source, "side": side}
    overview = render_overview_map(
        picture,
        areas,
        threats,
        starts,
        maps_folder / OVERVIEW_FILE,
        aspect=_OVERVIEW_ASPECT,
        width=_OVERVIEW,
        ground=ground,
        **tiles,
    )
    zooms = [
        render_objective_zoom(
            theatre,
            objective,
            area,
            threats,
            maps_folder / zoom_file(area.zone),
            width=_ZOOM,
            inset=_inset(overview, area),
            **tiles,
        )
        for objective, area in areas
    ]
    combat = (
        render_overview_map(
            picture,
            areas,
            threats,
            starts,
            maps_folder / COMBAT_AREA_FILE,
            aspect=_OVERVIEW_ASPECT,
            width=_OVERVIEW,
            combat_area=True,
            ground=ground,
            **tiles,
        )
        if areas or threats
        else None
    )
    if combat is None:  # the promotion takes the combat-area map first: an earlier run's would be stale
        (maps_folder / COMBAT_AREA_FILE).unlink(missing_ok=True)
    recon, recon_warnings = load_recon(recon_folder, [area.zone for _, area in areas])
    warnings += recon_warnings

    strategic = _prose_page(t("mission_briefing.title.strategic"), content.strategic_situation)
    situation = _situation_page(content, picture, starts)
    ato = _ato_page(content, picture, warnings)
    flow = _flow_page(content, picture, starts, areas, roles[1] if roles else set())
    navigation = navigation_page(theatre, picture)
    if navigation:
        navigation.title = t("mission_briefing.title.flight_plan")
    radios = radio_plans(table, side, templates=dynamic)
    frequencies = _frequencies_page(picture, radios, content.flights)
    deployment = _deployment_page(content)
    coordination = _coordination_page(content)
    steps = _step_sheets(content, picture, starts, overview, warnings)
    targets = _target_sheets(picture, areas, zooms, warnings)
    opening: list[Page | TablePage] = [*([strategic] if strategic else []), situation, ato]
    closing: list[Page | CardSection | TablePage] = [
        flow,
        *([deployment] if deployment else []),
        *([coordination] if coordination else []),
        *([navigation] if navigation else []),
        *([steps] if steps else []),
        frequencies,
        targets,
    ]
    titles = [
        *[page.title for page in opening],
        t("campaign.mission_deck.title.tactical"),
        *([t("mission_briefing.title.combat_area")] if combat else []),
        *[_objective_title(objective, area) for objective, area in areas],
        *[page.title for page in closing],
    ]
    numbered = [f"{n}. {title}" for n, title in enumerate(titles, 1)]
    for page, title in zip(opening, numbered, strict=False):
        page.title = title
    for item, title in zip(closing, numbered[-len(closing) :], strict=True):
        item.title = title
    maps_at = len(opening)

    with tempfile.TemporaryDirectory() as scratch:
        deck_pages, pdf_pages = _write_deck_and_pdf(
            folder,
            content,
            picture,
            numbered,
            opening,
            closing,
            maps_at,
            areas,
            overview,
            zooms,
            Path(scratch),
            combat,
            recon,
        )
    out = folder / DECK_FILE
    pdf = folder / PDF_FILE

    maps = [overview.path, *([combat.path] if combat else []), *(zoom.path for zoom in zooms)]
    pages = (
        _kneeboard_pages(folder, picture, areas, threats, starts, tiles, content, recon, ground) if kneeboard else {}
    )
    archives = built_archives(folder, miz)
    keys, other, moved = _write_pictures(folder, archives, maps, side, pages, kneeboard_replace and kneeboard)
    if moved:
        warnings.append(
            f"kneeboard_replace: {len(moved)} kneeboard page(s) not written by this tool taken out of the mission "
            f"({', '.join(moved)}); the folder's copies are in {_KNEEBOARD_BACKUP}/."
        )
    if other:
        warnings.append(
            f"{_PICTURE_KEYS['red' if side == 'blue' else 'blue']} holds {len(other)} picture(s): a player of no "
            "side sees red's then blue's (known limitation briefing-pictures-red-then-blue). Left as they are."
        )
    offline = overview.offline or bool(combat and combat.offline) or any(zoom.offline for zoom in zooms)
    if offline:
        warnings.append("Map tiles could not be downloaded: the maps stand on a plain background.")
    photos = [photo.image for photos in recon.values() for photo in photos]
    return BriefingReport(
        out, deck_pages, maps, keys, miz, offline, warnings, pdf, pdf_pages, archives, list(pages), photos
    )
