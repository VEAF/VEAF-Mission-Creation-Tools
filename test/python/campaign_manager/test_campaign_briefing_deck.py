"""The campaign's strategic briefing, as a PPTX (FEAT-CAMPAIGN-BRIEFING-DECK ticket 04)."""

from __future__ import annotations

import copy
import re
from pathlib import Path

from campaign_fixture import PROSE, VALID
from campaign_manager.briefing_deck import (
    DECK_FILE,
    MAP_FILE,
    Block,
    Page,
    _body_height,
    _body_width,
    block_height,
    campaign_deck,
    line_count,
    paginate,
)
from campaign_manager.briefing_prose import BriefingProse, parse_prose
from campaign_manager.campaign_manager import initial_state, parse_campaign
from campaign_manager.models import CampaignDefinition
from pptx import Presentation
from veaf_libs.i18n import language, t


def _campaign() -> CampaignDefinition:
    campaign, issues = parse_campaign(copy.deepcopy(VALID))
    assert campaign is not None, issues
    return campaign


def _prose(campaign: CampaignDefinition) -> BriefingProse:
    prose, issues = parse_prose(copy.deepcopy(PROSE), campaign.missions, 1)
    assert prose is not None, issues
    return prose


def _offline(url: str) -> bytes | None:
    return None


def _slides(path: Path) -> list[list[str]]:
    """Every slide's text, frame by frame."""
    return [
        [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        for slide in Presentation(str(path)).slides
    ]


def _deck(tmp_path: Path, with_prose: bool = True) -> list[list[str]]:
    campaign = _campaign()
    with language("fr"):
        report = campaign_deck(
            campaign,
            initial_state(campaign),
            _prose(campaign) if with_prose else None,
            tmp_path / "mission-01",
            cache_dir=tmp_path / "tiles",
            fetch=_offline,
        )
    assert report.path == tmp_path / "mission-01" / DECK_FILE
    assert (tmp_path / "mission-01" / MAP_FILE).is_file()
    return _slides(report.path)


class TestTheDeck:
    def test_the_pages_follow_the_situation_brief(self, tmp_path: Path) -> None:
        titles = [slide[0] for slide in _deck(tmp_path)[1:]]
        assert titles == [
            "1. Situation stratégique",
            "2. Situation militaire",
            "3. Carte stratégique",
            "4. Mission et intention",
            "5. Objectifs de la campagne",
            "6. Concept d'opération",
            "7. Règles d'engagement",
            "8. Mission 1 — La porte de Poti",
            "Annexe — Règles de la campagne",
        ]

    def test_the_cover_names_the_operation_and_lists_the_pages(self, tmp_path: Path) -> None:
        cover = "\n".join(_deck(tmp_path)[0])
        assert "KOLKHIDA" in cover
        assert "Caucasus — 8 missions — situation avant la mission 1" in cover
        assert "Annexe — Règles de la campagne" in cover

    def test_the_prose_stands_in_its_place(self, tmp_path: Path) -> None:
        slides = {slide[0]: "\n".join(slide[1:]) for slide in _deck(tmp_path)[1:]}
        assert "Une négociation s'ouvre." in slides["1. Situation stratégique"]
        assert (
            "– Mode d'action le plus probable : Tenir Senaki, puis reprendre l'offensive."
            in slides["2. Situation militaire"]
        )
        assert "– But : Briser l'offensive." in slides["4. Mission et intention"]
        assert "Phase 1 — la porte de Poti" in slides["6. Concept d'opération"]
        assert "1. Prendre Poti\nSécuriser le port." in slides["8. Mission 1 — La porte de Poti"]

    def test_the_facts_come_from_the_campaign(self, tmp_path: Path) -> None:
        slides = {slide[0]: "\n".join(slide[1:]) for slide in _deck(tmp_path)[1:]}
        military = slides["2. Situation militaire"]
        assert "– Kobuleti : base défendue" in military
        assert "Réserves : volume inconnu. Le renseignement estime qu'elles transitent par Gudauta depot." in military
        objectives = slides["5. Objectifs de la campagne"]
        assert "– prendre et tenir : Senaki" in objectives
        assert "– détruire : Gudauta depot" in objectives
        assert "Missions restantes, celle-ci comprise : 8." in objectives

    def test_no_figure_is_given_for_the_enemy(self, tmp_path: Path) -> None:
        military = _deck(tmp_path)[2][1]
        enemy = military.split("Forces amies")[0]
        assert not re.search(r"\d", enemy), enemy

    def test_the_annex_is_generated_from_the_rules(self, tmp_path: Path) -> None:
        annex = "\n".join(_deck(tmp_path)[-1])
        assert "pendant 2 minute(s)" in annex
        assert "2 blindé(s), 1 défense(s) aérienne(s) et 1 transport(s)" in annex
        assert "jusqu'à 4 unité(s) perdue(s)" in annex

    def test_without_prose_the_deck_is_the_generated_half(self, tmp_path: Path) -> None:
        slides = _deck(tmp_path, with_prose=False)
        assert [slide[0] for slide in slides[1:]] == [
            "1. Situation militaire",
            "2. Carte stratégique",
            "3. Objectifs de la campagne",
            "Annexe — Règles de la campagne",
        ]
        with language("fr"):
            assert t("campaign.deck.prose_missing") in "\n".join(slides[0])


class TestWhereAPageEnds:
    def test_a_long_paragraph_takes_more_lines(self) -> None:
        width = _body_width()
        assert line_count("Court.", 14, width) == 1
        assert line_count("mot " * 400, 14, width) > 5

    def test_a_page_too_full_is_continued_and_every_part_fits(self) -> None:
        paragraphs = tuple(f"Paragraphe {n} " + "texte de briefing " * 20 for n in range(30))
        with language("fr"):
            pages = paginate(Page("Situation", [Block("Contexte", paragraphs)]))
        assert len(pages) > 1
        assert pages[1].title == "Situation (suite)"
        for page in pages:
            assert sum(block_height(block, _body_width()) for block in page.blocks) <= _body_height(False)
        assert sum(len(block.paragraphs) for page in pages for block in page.blocks) == 30

    def test_a_page_that_fits_stays_one_page(self) -> None:
        assert len(paginate(Page("Situation", [Block("Contexte", ("Un paragraphe.",))]))) == 1
