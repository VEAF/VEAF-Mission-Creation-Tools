"""The tactical maps and the mission briefing deck (FEAT-CAMPAIGN-MISSION-BRIEFING tickets 03 and 04)."""

from __future__ import annotations

import copy
import io
import itertools
import re
from pathlib import Path

import yaml
from campaign_fixture import PROSE, VALID, built_mission, mission_template
from campaign_manager.briefing_prose import BriefingProse, parse_prose
from campaign_manager.campaign_manager import initial_state, parse_campaign
from campaign_manager.campaign_worker import CampaignWorker
from campaign_manager.mission_deck import MISSION_DECK_FILE, mission_deck, objective_zones
from campaign_manager.mission_picture import MissionPicture, read_mission_picture
from campaign_manager.models import CampaignDefinition
from campaign_manager.tactical_map import render_objective_map, render_tactical_map
from PIL import Image
from pptx import Presentation
from veaf_libs.coordinates import to_dms, xy_to_latlon
from veaf_libs.i18n import language
from veaf_libs.map_labels import LabelPlacer, intersects
from veaf_libs.mission_validator import WARNING

#: The coming mission's page: its tasks name two zones, in this order.
TASKS = {
    1: {
        "title": "La porte de Senaki",
        "tasks": [
            {"title": "Frapper Gudauta depot — priorité 1", "text": "Commencer l'attrition."},
            {"title": "Tenir le ciel", "text": "Couvrir les frappes."},
            {"title": "Reconnaître Senaki — opportunité", "text": "Localiser la défense aérienne."},
        ],
    }
}


def _white(url: str) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (256, 256), (255, 255, 255)).save(buffer, format="PNG")
    return buffer.getvalue()


def _offline(url: str) -> bytes | None:
    return None


def _campaign() -> CampaignDefinition:
    campaign, issues = parse_campaign(copy.deepcopy(VALID))
    assert campaign is not None, issues
    return campaign


def _prose(campaign: CampaignDefinition) -> BriefingProse:
    raw = copy.deepcopy(PROSE)
    raw["missions"] = copy.deepcopy(TASKS)
    prose, issues = parse_prose(raw, campaign.missions, 1)
    assert prose is not None, issues
    return prose


def _picture(tmp_path: Path) -> MissionPicture:
    return read_mission_picture(built_mission(tmp_path / "mission"), "blue", tmp_path / "mission")


class TestTheLabels:
    def test_a_label_goes_beside_its_anchor_when_free(self) -> None:
        placer = LabelPlacer(500, 500)
        assert placer.place((100, 100, 120, 120), (50, 20)) == (126, 100)

    def test_a_label_avoids_a_symbol_and_another_label(self) -> None:
        placer = LabelPlacer(500, 500)
        placer.reserve((120, 90, 200, 130))  # a symbol on the right
        first = placer.place((100, 100, 120, 120), (50, 20))
        second = placer.place((100, 100, 120, 120), (50, 20))
        boxes = [(*first, first[0] + 50, first[1] + 20), (*second, second[0] + 50, second[1] + 20)]
        assert not intersects(boxes[0], boxes[1])
        assert not any(intersects(box, (120, 90, 200, 130)) for box in boxes)

    def test_a_label_stays_inside_the_picture(self) -> None:
        placer = LabelPlacer(200, 200)
        x, y = placer.place((180, 90, 200, 110), (50, 20))
        assert 0 <= x and x + 50 <= 200 and 0 <= y and y + 20 <= 200


class TestTheMaps:
    def test_the_symbols_stand_at_their_positions(self, tmp_path: Path) -> None:
        campaign = _campaign()
        report = render_tactical_map(
            campaign,
            initial_state(campaign),
            _picture(tmp_path),
            tmp_path / "map.png",
            cache_dir=tmp_path,
            fetch=_white,
        )
        assert not report.offline
        image = Image.open(report.path).convert("RGB")
        red, _, blue = image.getpixel(tuple(int(v) for v in report.symbols["carrier"]))
        assert blue > red + 50  # the carrier's triangle, in the players' colour
        red, _, blue = image.getpixel(tuple(int(v) for v in report.symbols["Senaki"]))
        assert red > blue  # a red zone
        assert {"carrier", "bullseye", "Overlord 1-1", "Arco 1-1", "Texaco 1-1"} <= set(report.symbols)

    def test_no_two_labels_overlap(self, tmp_path: Path) -> None:
        campaign = _campaign()
        report = render_tactical_map(
            campaign,
            initial_state(campaign),
            _picture(tmp_path),
            tmp_path / "map.png",
            cache_dir=tmp_path,
            fetch=_white,
        )
        # zones, QRA, carrier, three support aircraft, bullseye, the flight plan's three numbered points
        assert len(report.labels) == 3 + 1 + 1 + 3 + 1 + 3
        for a, b in itertools.combinations(report.labels, 2):
            assert not intersects(a, b), (a, b)
        for label in report.labels:
            for name in ("carrier", "bullseye"):
                x, y = report.symbols[name]
                assert not intersects(label, (x - 5, y - 5, x + 5, y + 5)), name

    def test_the_flight_plans_points_are_numbered_where_they_stand(self, tmp_path: Path) -> None:
        campaign = _campaign()
        picture = _picture(tmp_path)
        report = render_tactical_map(
            campaign, initial_state(campaign), picture, tmp_path / "map.png", cache_dir=tmp_path, fetch=_white
        )
        # planes and helicopters share POTI and KHOBI: one numbered point each, the bullseye third
        assert [name for name in report.symbols if name.startswith("waypoint ")] == [
            "waypoint 1",
            "waypoint 2",
            "waypoint 3",
        ]
        assert report.symbols["waypoint 3"] == report.symbols["bullseye"]

    def test_the_maps_render_without_the_network(self, tmp_path: Path) -> None:
        campaign = _campaign()
        state = initial_state(campaign)
        tactical = render_tactical_map(
            campaign, state, _picture(tmp_path), tmp_path / "t.png", cache_dir=tmp_path, fetch=_offline
        )
        zoom = render_objective_map(
            campaign, state, campaign.zone("Senaki"), tmp_path / "z.png", cache_dir=tmp_path, fetch=_offline
        )
        assert tactical.offline and zoom.offline and tactical.path.is_file() and zoom.path.is_file()

    def test_a_zoom_centres_its_zone(self, tmp_path: Path) -> None:
        campaign = _campaign()
        zoom = render_objective_map(
            campaign,
            initial_state(campaign),
            campaign.zone("Senaki"),
            tmp_path / "z.png",
            cache_dir=tmp_path,
            fetch=_white,
        )
        image = Image.open(zoom.path).convert("RGB")
        x, y = zoom.symbols["Senaki"]
        assert abs(x - image.width / 2) < 5 and abs(y - image.height / 2) < 5
        red, _, blue = image.getpixel((int(x), int(y)))
        assert red > blue


class TestTheObjectives:
    def test_the_zones_the_tasks_name_in_their_order(self) -> None:
        campaign = _campaign()
        page = _prose(campaign).missions[1]
        assert [zone.name for zone in objective_zones(campaign, page)] == ["Gudauta depot", "Senaki"]

    def test_without_a_page_the_campaigns_objectives(self) -> None:
        campaign = _campaign()
        assert [zone.name for zone in objective_zones(campaign, None)] == ["Senaki", "Gudauta depot"]


def _deck(tmp_path: Path) -> list[list[str]]:
    campaign = _campaign()
    with language("fr"):
        report = mission_deck(
            campaign,
            initial_state(campaign),
            _prose(campaign),
            _picture(tmp_path),
            tmp_path / "mission-01",
            cache_dir=tmp_path,
            fetch=_offline,
        )
    assert report.path == tmp_path / "mission-01" / MISSION_DECK_FILE
    assert all(path.is_file() for path in report.maps)
    return [
        [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        for slide in Presentation(str(report.path)).slides
    ]


class TestTheDeck:
    def test_the_pages_follow_the_veaf_mission_briefing(self, tmp_path: Path) -> None:
        titles = [slide[0] for slide in _deck(tmp_path)[1:]]
        assert [title for title in titles if "(suite)" not in title] == [
            "1. Situation générale",
            "2. ATO",
            "3. Situation tactique",
            "4. Situation tactique — Gudauta depot",
            "5. Situation tactique — Senaki",
            "6. Déroulement mission",
            "7. Plan de fréquences",
            "8. Plan de navigation",
            "9. Coordonnées des objectifs",
        ]

    def test_the_cover_gives_the_mission_its_date_and_time(self, tmp_path: Path) -> None:
        cover = "\n".join(_deck(tmp_path)[0])
        assert "KOLKHIDA 1/8" in cover
        assert "La porte de Senaki" in cover
        assert "01/06/2016 — " in cover

    def test_the_ato_has_the_flights_and_no_template(self, tmp_path: Path) -> None:
        ato = "\n".join("\n".join(slide) for slide in _deck(tmp_path) if slide[0].startswith("2. ATO"))
        assert "Uzi 1 — Stennis Hornet, 4 × FA-18C_hornet" in ato
        assert "Colt 1 — Kobuleti Viper, 2 × F-16C_50" in ato
        assert "Template" not in ato and "veafSpawn" not in ato
        assert ato.count("Arco 1-1") == 1 and ato.count("Texaco 1-1") == 1

    def test_frequencies_are_split_uhf_then_vhf(self, tmp_path: Path) -> None:
        (page,) = [slide for slide in _deck(tmp_path) if slide[0] == "7. Plan de fréquences"]
        text = page[1]
        uhf, vhf = text.split("\nVHF\n")
        assert uhf.startswith("UHF\nGarde : 243.0 MHz")
        assert "127.5" not in uhf and "Stennis (tour) : 127.5 MHz, TACAN 74X, ICLS 1, Link 4 336.0 MHz" in vhf
        assert "Kobuleti : 262.0 MHz — TACAN 67X" in uhf and "Kobuleti : 133.0 MHz (FM 40.8 MHz)" in vhf

    def test_the_weather_and_the_wind_from_where_it_comes(self, tmp_path: Path) -> None:
        text = "\n".join("\n".join(slide) for slide in _deck(tmp_path))
        assert "Vent : sol 270° / 16 kt" in text
        assert "Nuages épars, base 8200 ft" in text

    def test_coordinates_are_zone_centres_never_at_sixty_seconds(self, tmp_path: Path) -> None:
        slides = _deck(tmp_path)
        (page,) = [slide for slide in slides if slide[0] == "9. Coordonnées des objectifs"]
        assert "Gudauta depot : N43°06'00.00\" E040°34'48.00\"" in page[1]
        assert "se relèvent en vol" in page[1]
        assert not re.search(r"60\.00\"", "\n".join("\n".join(slide) for slide in slides))

    def test_the_navigation_plan_planes_then_helicopters_bullseye_included(self, tmp_path: Path) -> None:
        (page,) = [slide for slide in _deck(tmp_path) if slide[0] == "8. Plan de navigation"]
        poti = to_dms(*xy_to_latlon("Caucasus", -295152.0, 617091.0))
        planes, helicopters = page[1].split("\nHélicoptères\n")
        assert planes.startswith("Avions\n")
        assert f"1. POTI : {poti} — 10\u00a0000 ft BARO" in planes
        assert "3. BULLSEYE : " in planes and "20\u00a0000 ft BARO" in planes
        assert f"1. POTI : {poti} — 500 ft AGL" in helicopters
        assert "2. KHOBI : " in helicopters

    def test_a_zoom_page_says_the_task_on_its_zone(self, tmp_path: Path) -> None:
        (page,) = [slide for slide in _deck(tmp_path) if slide[0] == "5. Situation tactique — Senaki"]
        assert "Reconnaître Senaki — opportunité\nLocaliser la défense aérienne." in page[1]


def test_dms_rounds_once_so_never_sixty_seconds() -> None:
    # 42°08'59.9999" rounds to 42°09'00.00", not 42°08'60.00"
    assert to_dms(42 + 8 / 60 + 59.9999 / 3600, 41.67) == "N42°09'00.00\" E041°40'12.00\""
    assert to_dms(-1.5, -0.25) == "S01°30'00.00\" W000°15'00.00\""


class TestTheWorker:
    def _folder(self, tmp_path: Path) -> CampaignWorker:
        folder = tmp_path / "campaign"
        folder.mkdir()
        (folder / "campaign.yaml").write_text(yaml.safe_dump(copy.deepcopy(VALID)), encoding="utf-8")
        mission_template(folder)
        worker = CampaignWorker(folder, tile_cache=tmp_path / "tiles", tile_fetch=_offline)
        worker.init()
        worker.next()
        return worker

    def test_without_a_built_mission_it_says_how_to_build_it(self, tmp_path: Path) -> None:
        worker = self._folder(tmp_path)
        issues, report = worker.briefing()
        assert report is not None and report.mission_deck is None
        assert any(issue.level == WARNING and "build" in issue.message for issue in issues)

    def test_a_corrupt_build_is_a_warning_and_the_campaign_briefing_still_comes(self, tmp_path: Path) -> None:
        worker = self._folder(tmp_path)
        (worker.mission_folder(1) / "mission" / "Half_written.miz").write_bytes(b"PK\x03\x04 cut short")
        issues, report = worker.briefing()
        assert report is not None and report.path.is_file() and report.mission_deck is None
        assert any(issue.level == WARNING for issue in issues)
        _, next_report = worker.next()  # a refresh with the same file does not fail either
        assert next_report is not None

    def test_with_a_built_mission_both_briefings_are_written(self, tmp_path: Path) -> None:
        worker = self._folder(tmp_path)
        built_mission(worker.mission_folder(1) / "mission")
        issues, report = worker.briefing()
        assert report is not None, issues
        assert report.mission_deck == worker.mission_folder(1) / MISSION_DECK_FILE
        assert report.mission_deck.is_file()
        assert not any("build" in issue.message for issue in issues)
