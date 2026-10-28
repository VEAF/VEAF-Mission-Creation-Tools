"""The briefing at the level of "Marteau du Golan"'s hand-made one (FIX-BRIEFING-PARITY-MARTEAU).

A regeneration of the mission's briefing (2026-10-28) found nine defects; one class per defect, D1 being
in `test_dcs_recon.py`.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from pathlib import Path
from typing import Any

import pytest
from campaign_manager.mission_picture import FlightPlan, Support, Waypoint, read_mission_picture
from mission_briefing import briefing_pdf
from mission_briefing.briefing_maps import render_overview_map
from mission_briefing.briefing_pdf import Cover, write_pdf
from mission_briefing.briefing_tables import Table, TablePage, paginate_tables, timeline
from mission_briefing.mission_briefing_worker import (
    _ato_plans,
    _named_plan,
    _objective_blocks,
    _step_sheets,
    bullseye_text,
    build_mission_briefing,
    flies,
    objective_notes,
    objective_waypoint,
)
from mission_briefing.mission_reader import ground_picture, objective_area, operation_zones
from mission_briefing.models import (
    BriefingContent,
    GroundPicture,
    MapLine,
    Objective,
    ObjectiveArea,
    PackageFlight,
    RadioPlan,
    Threat,
)
from mission_tools.miz_tools import read_miz, write_miz
from objective_mission_fixture import KHOBI_ZONE, objective_mission
from PIL import Image
from pptx import Presentation
from veaf_libs.i18n import language
from veaf_mission_mcp.build_briefing import _plan_notes


def _white(url: str) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (256, 256), (255, 255, 255)).save(buffer, format="PNG")
    return buffer.getvalue()


def _unit(name: str, kind: str, x: float = 0.0, y: float = 0.0) -> dict[str, Any]:
    return {"name": name, "type": kind, "x": x, "y": y}


def _group(name: str, *units: dict[str, Any], **extra: Any) -> dict[str, Any]:
    return {"name": name, "units": dict(enumerate(units, 1)), **extra}


def _table(red: list[dict[str, Any]], blue: list[dict[str, Any]] | None = None, **extra: Any) -> dict[str, Any]:
    def side(groups: list[dict[str, Any]]) -> dict[str, Any]:
        return {"country": {1: {"vehicle": {"group": dict(enumerate(groups, 1))}}}}

    return {
        "coalition": {"red": side(red), "blue": side(blue or [])},
        "triggers": {"zones": {1: {"name": "SA6Sasa", "x": 0.0, "y": 0.0, "radius": 1000.0}}},
        **extra,
    }


class TestUnitTargets:
    """D3: the SA-6 was one sheet named ``battery``; it is a radar and four launchers."""

    def test_units_named_by_the_author_are_one_target_each_with_their_role(self) -> None:
        battery = _group(
            "SA6Sasa-battery",
            _unit("SA6Sasa radar", "Kub 1S91 str"),
            *(_unit(f"SA6Sasa lanceur {n}", "Kub 2P25 ln", n * 10.0) for n in range(1, 5)),
        )
        with language("fr"):
            targets = objective_area(_table([battery]), "SA6Sasa", "red").targets
        assert [target.name for target in targets] == ["Radar", "Lanceur 1", "Lanceur 2", "Lanceur 3", "Lanceur 4"]
        assert targets[0].kind == "Kub 1S91 str" and targets[0].role == "radar de la batterie"
        assert targets[1].role == "lanceur de la batterie"

    def test_a_group_of_unnamed_units_is_named_by_its_type_not_its_slug(self) -> None:
        tanks = _group("SA6Sasa-t72", *(_unit(f"SA6Sasa-t72 Unit #00{n}", "T-72B3") for n in range(1, 5)))
        guns = _group("SA6Sasa-2s1", _unit("SA6Sasa-2s1-1", "SAU Gvozdika"), _unit("SA6Sasa-2s1-2", "SAU Gvozdika"))
        with language("fr"):
            targets = objective_area(_table([tanks, guns]), "SA6Sasa", "red").targets
        assert [(target.name, target.kind, target.role) for target in targets] == [
            ("T-72B3", "4 × T-72B3", "char"),
            ("SAU Gvozdika", "2 × SAU Gvozdika", "artillerie"),
        ]

    def test_a_veaf_spawn_command_is_named_by_its_family_and_numbered(self) -> None:
        groups = [
            _group("SA6Sasa-armor-1", _unit('#command="-armor" SA6Sasa-armor-1', "T-72B")),
            _group("SA6Sasa-armor-2", _unit('#command="-armor" SA6Sasa-armor-2', "T-72B")),
            _group("SA6Sasa-sa13", _unit('#command="-sa13_squad" SA6Sasa-sa13', "Strela-10M3")),
        ]
        with language("fr"):
            targets = objective_area(_table(groups), "SA6Sasa", "red").targets
        assert [target.name for target in targets] == ["Blindés 1", "Blindés 2", "SA-13"]
        assert all(target.spawned for target in targets)
        assert "varie d'une partie à l'autre" in targets[0].kind

    def test_the_target_sheets_carry_the_role_and_the_bullseye(self, tmp_path: Path) -> None:
        objective_mission(tmp_path)
        content = BriefingContent(title="T", objectives=(Objective("Khobi", "Le dépôt"),))
        report = build_mission_briefing(tmp_path, content, cache_dir=tmp_path / "tiles", fetch=_white)
        text = "\n".join(
            shape.text_frame.text
            for slide in Presentation(str(report.deck)).slides
            for shape in slide.shapes
            if shape.has_text_frame
        )
        assert "missiles sol-air à courte portée" in text  # the SA-8's role
        assert "Bullseye : " in text and " nm" in text


class TestBullseyeAndFlights:
    """D4: no bullseye on a target or a step, no "Pour : <flight>" on an objective."""

    def test_the_bullseye_bearing_and_range_on_the_grid_the_legs_use(self, tmp_path: Path) -> None:
        picture = read_mission_picture(objective_mission(tmp_path), "blue", tmp_path)
        bx, by = picture.bullseye
        with language("fr"):
            assert bullseye_text(picture, bx + 18520, by) == "Bullseye : 000° / 10.0 nm"  # x is north
            assert bullseye_text(picture, bx, by + 1852) == "Bullseye : 090° / 1.0 nm"
            assert bullseye_text(picture, bx, by) == "Bullseye : 000° / 0.0 nm"

    def test_an_objective_says_for_whom_its_bullseye_and_its_point(self, tmp_path: Path) -> None:
        picture = read_mission_picture(objective_mission(tmp_path), "blue", tmp_path)
        area = ObjectiveArea("Khobi", KHOBI_ZONE[1], KHOBI_ZONE[2], KHOBI_ZONE[3], label="KHOBI")
        objective = Objective("Khobi", "Le dépôt", ("Le détruire.",), flights=("Ninja", "Arctic"))
        with language("fr"):
            notes = [text for text, _ in objective_notes(objective, area, picture)]
        assert notes[:2] == ["Le dépôt", "Le détruire."]
        assert "Pour : Ninja et Arctic" in notes
        assert any(note.startswith("Bullseye : ") for note in notes)
        assert "Point : KHOBI du plan de vol" in notes

    def test_the_point_on_an_objective_is_the_flight_plans_nearest_within_its_zone(self, tmp_path: Path) -> None:
        picture = read_mission_picture(objective_mission(tmp_path), "blue", tmp_path)
        name, x, y, radius = KHOBI_ZONE
        assert objective_waypoint(ObjectiveArea(name, x, y, radius), picture) == "KHOBI"
        assert objective_waypoint(ObjectiveArea("Far", x + 90_000, y, 1000), picture) == ""


class TestMainAndSecondary:
    """D5: every objective was "main", the counter-attack included."""

    @staticmethod
    def _yaml(folder: Path, operation: bool = True) -> None:
        zones = [
            "      - type: zone\n        zone_name: Khan\n        chained_zones:\n          - Counter\n",
            "      - type: zone\n        zone_name: Sa6\n",
            "      - type: zone\n        zone_name: Counter\n",
        ]
        if operation:
            zones.append("      - type: operation\n        zone_name: Op\n        tasking_orders:\n          - zone_name: Khan\n")
        (folder / "mission.yaml").write_text(
            "modules:\n  COMBATZONE:\n    enabled: true\n    combat_zones:\n" + "".join(zones), encoding="utf-8"
        )

    def test_the_operation_tasks_and_the_chained_zones_are_read(self, tmp_path: Path) -> None:
        self._yaml(tmp_path)
        assert operation_zones(tmp_path) == ({"khan"}, {"counter"})
        self._yaml(tmp_path, operation=False)
        assert operation_zones(tmp_path) is None

    def test_main_then_secondary_and_no_chained_zone(self) -> None:
        areas = [
            (Objective("Khan", "La ligne"), ObjectiveArea("Khan", 0, 0, 1)),
            (Objective("Sa6", "Le SA-6"), ObjectiveArea("Sa6", 0, 0, 1, secondary=True)),
            (Objective("Counter", "La contre-attaque"), ObjectiveArea("Counter", 0, 0, 1, secondary=True)),
        ]
        with language("fr"):
            blocks = _objective_blocks(areas, {"counter"})
        assert [(block.heading, block.paragraphs) for block in blocks] == [
            ("Objectifs principaux", ("1. La ligne",)),
            ("Objectifs secondaires", ("– Le SA-6",)),
        ]

    def test_without_an_operation_every_objective_is_listed(self) -> None:
        areas = [(Objective("A", "Un"), ObjectiveArea("A", 0, 0, 1)), (Objective("B", "Deux"), ObjectiveArea("B", 0, 0, 1))]
        with language("fr"):
            (block,) = _objective_blocks(areas, set())
        assert block.paragraphs == ("1. Un", "2. Deux")


class TestFrequencyPlans:
    """D6: seven pages of radios, one per dynamic-slot type, and the 61 types on the flight card."""

    def test_an_atos_aircraft_names_the_dcs_types_it_flies(self) -> None:
        assert flies("F-16C", "F-16C_50") and flies("F/A-18C", "FA-18C_hornet") and flies("A-10C", "A-10C_2")
        assert flies("AH-64D, OH-58D, Mi-24P", "OH58D") and flies("AV-8B", "AV8BNA")
        assert not flies("F-16C", "F-15C") and not flies("A-10C", "AV8BNA")

    def test_one_block_per_radio_plan_of_the_atos_aircraft(self) -> None:
        package = ((1, "Package", 291.0), (2, "Magic", 282.2))
        bare = ((1, "", 30.0), (2, "", 31.0))
        plans = [
            RadioPlan(1, ("F-16C_50", "FA-18C_hornet", "Su-25T"), package),
            RadioPlan(2, ("AH-64D_BLK_II",), package),
            RadioPlan(3, ("AH-64D_BLK_II",), bare),
            RadioPlan(1, ("MiG-29S",), ((1, "Other", 250.0),)),
        ]
        flights = (PackageFlight("Ninja", aircraft="F-16C"), PackageFlight("Blade", aircraft="AH-64D"))
        (plan,) = _ato_plans(plans, flights) or []
        assert plan.radios == (1, 2)
        assert plan.flights == ("Ninja", "Blade")  # named by the flights, in the ATO's order
        assert plan.channels == package

    def test_no_flight_no_filter(self) -> None:
        plans = [RadioPlan(1, ("F-16C_50",), ((1, "Package", 291.0),))]
        assert _ato_plans(plans, ()) is None
        assert _ato_plans(plans, (PackageFlight("Ninja", aircraft="Spitfire"),)) is None

    def test_a_dynamic_plan_is_named_by_the_flights_that_fly_it(self) -> None:
        plan = FlightPlan("plane", ("A-10C_2", "F-16C_50", "MiG-29S", "Su-25T"), ())
        flights = (PackageFlight("Ninja", aircraft="F-16C"), PackageFlight("Blade", aircraft="AH-64D"))
        assert _named_plan(plan, flights).aircraft == ("Ninja",)
        assert _named_plan(plan, (PackageFlight("X", aircraft="Spitfire"),)).aircraft == plan.aircraft


class TestPdfAndTables:
    """D7: a justified PDF, the ATO, the deployment and the code words as plain lines."""

    def test_the_pdf_aligns_left(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        aligns: list[Any] = []
        original = briefing_pdf._Pdf.multi_cell

        def spy(self: Any, *args: Any, **kwargs: Any) -> Any:
            aligns.append(kwargs.get("align"))
            return original(self, *args, **kwargs)

        monkeypatch.setattr(briefing_pdf._Pdf, "multi_cell", spy)
        from campaign_manager.briefing_deck import Block, Page

        cover = Cover("T", "S", "D", "Sommaire", ["1. P"])
        write_pdf(tmp_path / "b.pdf", cover, [Page("1. P", [Block("H", ("Un long paragraphe " * 30,))])])
        assert aligns and set(aligns) == {"L"}

    def test_the_deployment_splits_into_when_who_what(self) -> None:
        rows, prose = timeline(
            (
                "H+0 — Ninja, Arctic : mise en route.",
                "H+15 — Les QRA de Mezzeh et de Khalkhalah sont en alerte : elles décollent.",
                "« MARTEAU » — Mission Maker : la poussée.",
                "Qui parle à qui. Package (291.0) : tout le monde.",
            )
        )
        assert rows == [
            ("H+0", "Ninja, Arctic", "mise en route."),
            ("H+15", "", "Les QRA de Mezzeh et de Khalkhalah sont en alerte : elles décollent."),
            ("« MARTEAU »", "Mission Maker", "la poussée."),
        ]
        assert prose == ["Qui parle à qui. Package (291.0) : tout le monde."]

    def test_a_long_table_runs_on_with_its_header(self) -> None:
        rows = tuple((f"H+{n}", "Ninja", "Une ligne de texte assez longue " * 3) for n in range(40))
        pages = paginate_tables(TablePage("T", [Table(None, ("Quand", "Qui", "Quoi"), (0.2, 0.2, 0.6), rows)]))
        assert len(pages) > 1
        assert sum(len(table.rows) for page in pages for table in page.tables) == 40

    def test_the_ato_deployment_and_code_words_are_tables_in_the_deck_and_the_pdf(self, tmp_path: Path) -> None:
        from mission_briefing.models import CodeWord

        objective_mission(tmp_path)
        content = BriefingContent(
            title="T",
            deployment=("H+0 — Ninja : décollage.", "Coordination sur Package."),
            code_words=(CodeWord("PORTE OUVERTE", "Le SA-6 est détruit.", "Blade passe."),),
        )
        report = build_mission_briefing(tmp_path, content, cache_dir=tmp_path / "tiles", fetch=_white)
        tables = {
            slide.shapes[0].text_frame.text.split(". ", 1)[-1]
            for slide in Presentation(str(report.deck)).slides
            if any(shape.has_table for shape in slide.shapes)
        }
        assert {"ATO", "Déploiement et coordination", "Mots-codes"} <= tables
        assert report.pdf is not None and report.pdf.is_file()


class TestNotesPerPlan:
    """D8: jets and helicopters share SASA, KHAN… with different roles, and a note was one per name."""

    def test_a_plans_note_wins_over_the_shared_one(self, tmp_path: Path) -> None:
        picture = read_mission_picture(objective_mission(tmp_path), "blue", tmp_path)
        point = Waypoint("KHAN", 1000.0, 2000.0, 3000.0, "BARO")
        picture = replace(
            picture, flight_plans=[FlightPlan("plane", ("F-16C_50",), (point,)), FlightPlan("helicopter", ("AH-64D",), (point,))]
        )
        content = BriefingContent(
            title="T", waypoint_notes={"KHAN": "Au contact."}, plan_notes={"helicopter": {"KHAN": "Retour par la route."}}
        )
        warnings: list[str] = []
        overview = render_overview_map(picture, [], [], [], tmp_path / "o.png", cache_dir=tmp_path / "tiles", fetch=_white)
        with language("fr"):
            section = _step_sheets(content, picture, [], overview, warnings)
        assert section is not None
        (plane,), (helicopter,) = (cards for _, cards in section.groups)
        assert "Au contact." in plane.lines and "Retour par la route." not in plane.lines
        assert "Retour par la route." in helicopter.lines and "Au contact." not in helicopter.lines
        assert not warnings

    def test_the_parameter_takes_a_map_per_plan(self) -> None:
        notes = {"PUSH": "Attente.", "helicopter": {"KHAN": "Retour."}, "jets": {"KHAN": "Contact."}}
        assert _plan_notes(notes) == {"helicopter": {"KHAN": "Retour."}, "plane": {"KHAN": "Contact."}}
        with pytest.raises(ValueError, match="plan"):
            _plan_notes({"boats": {"KHAN": "?"}})


class TestKneeboardReplace:
    """D9: the hand-made kneeboard stayed beside the tool's, 20 pages and 24."""

    @staticmethod
    def _hand_made(folder: Path, miz: Path) -> Path:
        pages = folder / "src" / "mission" / "KNEEBOARD" / "IMAGES"
        pages.mkdir(parents=True, exist_ok=True)
        (pages / "01-sitac-generale.jpg").write_bytes(b"hand")
        (pages / "notes.txt").write_text("not a page", encoding="utf-8")
        mission = read_miz(miz)
        write_miz(mission, miz, additional_files={"KNEEBOARD/IMAGES/01-sitac-generale.jpg": b"hand"})
        return pages

    def test_left_in_place_unless_asked(self, tmp_path: Path) -> None:
        miz = objective_mission(tmp_path)
        pages = self._hand_made(tmp_path, miz)
        build_mission_briefing(tmp_path, BriefingContent(title="T"), cache_dir=tmp_path / "tiles", fetch=_white)
        assert (pages / "01-sitac-generale.jpg").is_file()
        assert "KNEEBOARD/IMAGES/01-sitac-generale.jpg" in zipfile.ZipFile(miz).namelist()

    def test_replaced_when_asked_and_kept_in_a_backup(self, tmp_path: Path) -> None:
        miz = objective_mission(tmp_path)
        pages = self._hand_made(tmp_path, miz)
        report = build_mission_briefing(
            tmp_path, BriefingContent(title="T"), cache_dir=tmp_path / "tiles", fetch=_white, kneeboard_replace=True
        )
        assert not (pages / "01-sitac-generale.jpg").exists()
        assert (pages / "notes.txt").is_file()  # not a picture: not the tool's to move
        assert list(tmp_path.glob(".veaf-backups/kneeboard-*/01-sitac-generale.jpg"))
        names = zipfile.ZipFile(miz).namelist()
        assert "KNEEBOARD/IMAGES/01-sitac-generale.jpg" not in names
        assert any(name.startswith("KNEEBOARD/IMAGES/01-briefing-") for name in names)
        assert any("01-sitac-generale.jpg" in warning for warning in report.warnings)


class TestTheaterMaps:
    """D2: the overview framed on the objectives, filled pink, trigger names; the combat map bare."""

    @staticmethod
    def _picture(tmp_path: Path) -> Any:
        picture = read_mission_picture(objective_mission(tmp_path), "blue", tmp_path)
        _, x, y, _ = KHOBI_ZONE
        helicopter = FlightPlan("helicopter", ("AH-64D_BLK_II",), (Waypoint("MASADA", x - 8000, y - 4000, 300, "RADIO"),))
        tanker = Support("Texaco 1-1", "tanker", "KC-135", 290.1, None, ((x - 60_000, y - 30_000), (x - 70_000, y - 30_000)))
        return replace(picture, flight_plans=[*picture.flight_plans, helicopter], support=[*picture.support, tanker])

    @staticmethod
    def _ground() -> GroundPicture:
        _, x, y, _ = KHOBI_ZONE
        return GroundPicture(
            friendlies=((x - 5000, y - 5000),),
            jtacs=(("JTAC AbuNada", x - 4000, y - 6000),),
            axes=(("red", ((x, y), (x - 3000, y - 3000))),),
            lines=(MapLine("Ligne de contact", (204, 0, 0), 4, False, ((x - 6000, y - 9000), (x - 6000, y + 3000))),),
        )

    def test_the_overview_frames_the_theatre_and_names_as_the_flight_plan(self, tmp_path: Path) -> None:
        picture = self._picture(tmp_path)
        name, x, y, radius = KHOBI_ZONE
        area = ObjectiveArea(name, x, y, radius, label="KHOBI")
        threat = (Threat("SA-6 Gori", "SA-6 Site", "SA-6", engagement_nm=13), (x + 20_000, y))
        with language("fr"):
            report = render_overview_map(
                picture,
                [(Objective(name, "Le dépôt"), area)],
                [threat],
                [],
                tmp_path / "o.png",
                cache_dir=tmp_path / "tiles",
                fetch=_white,
                width=1600,
                aspect=2.0,
                ground=self._ground(),
            )
        assert report.locate is not None
        width, height = Image.open(report.path).size
        tx, ty = report.locate(x - 65_000, y - 30_000)  # the tanker's track, far from the route
        assert 0 <= tx <= width and 0 <= ty <= height
        assert {"key helicopter_route", "key line 0", "key route"} <= set(report.symbols)
        assert "key friendly" not in report.symbols  # the overview leaves the ground to the combat map

    def test_a_long_range_reach_is_a_tint_not_a_fill(self, tmp_path: Path) -> None:
        picture = replace(self._picture(tmp_path), qra_zones=[], flight_plans=[], support=[], carriers=[])
        _, x, y, _ = KHOBI_ZONE
        threat = (Threat("SA-11", "SA-6 Site", "SA-11", engagement_nm=25), (x, y))
        report = render_overview_map(picture, [], [threat], [], tmp_path / "o.png", cache_dir=tmp_path / "tiles", fetch=_white)
        assert report.locate is not None
        px, py = report.locate(x + 6000, y + 6000)  # inside the circle, away from its centre and its label
        red, green, _ = Image.open(report.path).convert("RGB").getpixel((round(px), round(py)))
        assert red - green < 30  # a filled circle (alpha 70) took green down by some 60

    def test_the_combat_map_shows_the_friendlies_the_jtac_the_axes_and_the_ranges(self, tmp_path: Path) -> None:
        picture = self._picture(tmp_path)
        name, x, y, radius = KHOBI_ZONE
        area = ObjectiveArea(name, x, y, radius, label="KHOBI")
        threat = (Threat("SA-8", "Khobi-SA8", "SA-8", engagement_nm=5), (x + 300, y + 200))
        with language("fr"):
            report = render_overview_map(
                picture,
                [(Objective(name, "Le dépôt"), area)],
                [threat],
                [],
                tmp_path / "c.png",
                cache_dir=tmp_path / "tiles",
                fetch=_white,
                combat_area=True,
                ground=self._ground(),
            )
        assert {"key friendly", "key jtac", "key axis_enemy", "JTAC AbuNada"} <= set(report.symbols)

    def test_the_ground_is_read_from_the_mission(self) -> None:
        table = _table(
            red=[_group("Counter", _unit("c", "T-72B"), route={"points": {1: {"x": 0, "y": 0}, 2: {"x": 5, "y": 5}}})],
            blue=[_group("JTAC-AbuNada", _unit("j", "Hummer", 1, 2)), _group("Brigade", _unit("b", "Merkava_Mk4", 3, 4))],
            drawings={
                "layers": {
                    1: {
                        "name": "Blue",
                        "objects": {
                            1: {
                                "name": "Ligne de contact",
                                "primitiveType": "Line",
                                "colorString": "0xCC0000FF",
                                "thickness": 4,
                                "style": "boundry1",
                                "mapX": 100.0,
                                "mapY": 200.0,
                                "points": {1: {"x": 0.0, "y": 0.0}, 2: {"x": 10.0, "y": -5.0}},
                            }
                        },
                    },
                    2: {"name": "Red", "objects": {1: {"name": "Theirs", "primitiveType": "Line", "points": {}}}},
                }
            },
        )
        ground = ground_picture(table, "blue")
        assert ground.friendlies == ((3.0, 4.0),)
        assert ground.jtacs == (("JTAC AbuNada", 1.0, 2.0),)
        assert ground.axes == (("red", ((0.0, 0.0), (5.0, 5.0))),)
        assert ground.lines == (MapLine("Ligne de contact", (204, 0, 0), 4, False, ((100.0, 200.0), (110.0, 195.0))),)
