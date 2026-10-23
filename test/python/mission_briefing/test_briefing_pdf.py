"""The briefing's PDF, written by the executable (FIX-DEMO-VIDEO-FINDINGS 08).

The PDF is read back with `pypdf`: what the test compares is the text a reader would copy out of it, accents
included, against the deck's own words.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pytest
from mission_briefing.briefing_pdf import Cover, MapSheet, write_pdf
from mission_briefing.mission_briefing_worker import build_mission_briefing
from mission_briefing.models import BriefingContent, Objective, Threat
from objective_mission_fixture import objective_mission
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

CONTENT = BriefingContent(
    title="Marteau de Khobi — le nœud",
    context=("Les dépôts de Khobi ravitaillent le front : ça brûle, où l'été ne finit pas.",),
    mission=("Détruire le dépôt", "Rentrer à Kobuleti"),
    objectives=(Objective("Khobi", "Détruire le dépôt de Khobi", ("Le dépôt et sa défense SA-8.",)),),
    threats=(Threat("SA-6 Gori", "SA-6 Site", "SA-6", engagement_nm=13, detection_nm=30, source="mesuré en jeu"),),
    air_opposition=("Alerte d'interception à Senaki.",),
    other=("Attaque en piqué conseillée.",),
)


def _white(url: str) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (256, 256), (255, 255, 255)).save(buffer, format="PNG")
    return buffer.getvalue()


def _words(text: str) -> set[str]:
    return set(re.findall(r"[^\s]+", text))


def _pdf_text(pdf: Path) -> str:
    return "\n".join(page.extract_text() for page in PdfReader(str(pdf)).pages)


def _deck_text(deck: Path) -> str:
    return "\n".join(
        shape.text_frame.text
        for slide in Presentation(str(deck)).slides
        for shape in slide.shapes
        if shape.has_text_frame
    )


@pytest.fixture
def report(tmp_path: Path):
    objective_mission(tmp_path)
    return build_mission_briefing(tmp_path, CONTENT, cache_dir=tmp_path / "tiles", fetch=_white)


def test_the_briefing_gives_a_pdf_beside_the_deck(report, tmp_path: Path) -> None:
    assert report.pdf == tmp_path / "docs" / "briefing.pdf"
    assert report.pdf.read_bytes().startswith(b"%PDF")
    assert report.pdf_pages >= report.pages
    assert len(PdfReader(str(report.pdf)).pages) == report.pdf_pages


def test_what_a_reader_copies_out_of_the_pdf_is_the_decks_words_accents_included(report) -> None:
    pdf, deck = _pdf_text(report.pdf), _deck_text(report.deck)
    missing = sorted(_words(deck) - _words(pdf))
    assert missing == [], f"words of the deck the PDF does not carry: {missing[:10]}"
    for accented in ("dépôt", "Détruire", "nœud", "été", "défense"):
        assert accented in pdf, accented


def test_the_font_is_the_bundled_one_not_a_core_font(report) -> None:
    assert b"DejaVu" in report.pdf.read_bytes()


def test_a_page_with_a_map_carries_its_image(tmp_path: Path) -> None:
    image = tmp_path / "map.png"
    Image.new("RGB", (300, 200), (200, 220, 200)).save(image)
    out = tmp_path / "x.pdf"
    pages = write_pdf(
        out,
        Cover("Titre", "Sous-titre", "13 octobre 2015", "Sommaire", ["1. Carte"]),
        [MapSheet("1. Carte", image, [("Zone", True), ("Détruire le nœud.", False)])],
    )
    reader = PdfReader(str(out))
    assert pages == len(reader.pages) == 2
    assert len(reader.pages[1].images) == 1
    assert "Détruire le nœud." in reader.pages[1].extract_text()


def test_a_very_long_paragraph_runs_on_to_a_new_page_instead_of_overflowing(tmp_path: Path) -> None:
    from campaign_manager.briefing_deck import Block, Page

    out = tmp_path / "long.pdf"
    page = Page(
        "1. Long", [Block("Contexte", tuple(f"Ligne {n} : les dépôts ravitaillent le front." for n in range(120)))]
    )
    pages = write_pdf(out, Cover("T", "S", "D", "Sommaire", []), [page])
    assert pages > 2
    assert "Ligne 119" in _pdf_text(out)
