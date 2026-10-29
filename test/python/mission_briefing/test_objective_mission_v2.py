"""What "Marteau du Golan" did by hand, in the briefing (FEAT-OBJECTIVE-MISSION-V2 02 to 04).

A combat-area map beside the overview, kneeboard maps turned attack axis up, the code words, and the
objectives' recon photos read from JPEG and their JSON sidecars.
"""

from __future__ import annotations

import io
import json
import math
import re
from datetime import date
from pathlib import Path
from typing import NamedTuple

import pytest
from campaign_manager.mission_picture import read_mission_picture
from geographiclib.geodesic import Geodesic  # type: ignore[import-untyped]
from mission_briefing.briefing_maps import attack_axis_rotation, render_overview_map
from mission_briefing.briefing_recon import caption, load_recon
from mission_briefing.mission_briefing_worker import PICTURE_PREFIX, build_mission_briefing
from mission_briefing.mission_reader import departures, objective_area
from mission_briefing.models import BriefingContent, BriefingReport, CodeWord, Objective, ReconPhoto
from mission_tools.miz_tools import read_miz
from objective_mission_fixture import KHOBI_ZONE, objective_mission
from offline import offline
from PIL import Image
from pptx import Presentation
from veaf_libs.coordinates import xy_to_latlon
from veaf_libs.magnetic import declination

CONTENT = BriefingContent(
    title="Marteau de Khobi",
    objectives=(Objective("Khobi", "Détruire le dépôt de Khobi", ("Le dépôt et sa défense SA-8.",)),),
)


def _white(url: str) -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (256, 256), (255, 255, 255)).save(buffer, format="PNG")
    return buffer.getvalue()


def _inputs(tmp_path: Path) -> dict:
    miz = objective_mission(tmp_path / "mission")
    content = read_miz(miz).mission_content or {}
    picture = read_mission_picture(miz, "blue", tmp_path / "mission")
    return {
        "picture": picture,
        "objectives": [(CONTENT.objectives[0], objective_area(content, "Khobi", "red"))],
        "threats": [],
        "departures": departures(content, "blue", picture.theatre),
    }


CODES = (
    CodeWord("PORTE OUVERTE", "Le SA-6 de Sa'sa' est détruit.", "Blade franchit le col."),
    CodeWord("MARTEAU", "Khan Arnabah est tombé."),
)


def _build(folder: Path, content: BriefingContent = CONTENT, **extra):
    return build_mission_briefing(folder, content, cache_dir=folder / "tiles", fetch=_white, **extra)


class Built(NamedTuple):
    folder: Path
    report: BriefingReport


@pytest.fixture(scope="module")
def built(tmp_path_factory: pytest.TempPathFactory) -> Built:
    """`CONTENT`'s briefing, built once for the tests that only read it: a build takes seconds."""
    folder = tmp_path_factory.mktemp("briefing")
    with offline(folder):
        objective_mission(folder)
        return Built(folder, _build(folder))


def _slides(deck: Path) -> list[list[str]]:
    """Each slide's texts: its text boxes, then each cell of its tables."""
    return [
        [shape.text_frame.text for shape in slide.shapes if shape.has_text_frame]
        + [cell.text for shape in slide.shapes if shape.has_table for row in shape.table.rows for cell in row.cells]
        for slide in Presentation(str(deck)).slides
    ]


def _photo(folder: Path, stem: str, sidecar: dict | None, size: tuple[int, int] = (1600, 900)) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    image = folder / f"{stem}.jpg"
    Image.new("RGB", size, (90, 110, 70)).save(image, format="JPEG")
    if sidecar is not None:
        (folder / f"{stem}.json").write_text(json.dumps(sidecar, ensure_ascii=False), encoding="utf-8")
    return image


class TestTheAttackAxis:
    @pytest.mark.parametrize(
        ("target", "turn"), [((1, 0), 0), ((0, 1), 90), ((-1, 0), 180), ((0, -1), -90), ((0.01, 0.01), 45)]
    )
    def test_the_axis_from_the_departure_to_the_objectives_points_up(self, target, turn) -> None:
        assert attack_axis_rotation((0, 0), [target]) == pytest.approx(turn)

    def test_no_axis_without_an_origin_or_a_target(self) -> None:
        assert attack_axis_rotation(None, [(1, 2)]) is None
        assert attack_axis_rotation((0, 0), []) is None
        assert attack_axis_rotation((5, 5), [(5, 5)]) is None

    def test_turned_by_it_the_objective_stands_above_the_departure(self, tmp_path: Path) -> None:
        inputs = _inputs(tmp_path)
        _, dx, dy = inputs["departures"][0]
        area = inputs["objectives"][0][1]
        theatre = inputs["picture"].theatre
        turn = attack_axis_rotation(xy_to_latlon(theatre, dx, dy), [xy_to_latlon(theatre, area.x, area.y)])
        report = render_overview_map(
            **inputs, out=tmp_path / "kb.png", cache_dir=tmp_path, fetch=_white, aspect=2 / 3, width=800, rotation=turn
        )
        assert report.locate is not None
        ox, oy = report.locate(area.x, area.y)
        sx, sy = report.locate(dx, dy)
        # up the page, and along it: the axis is within a few degrees of vertical
        assert oy < sy
        assert abs(math.degrees(math.atan2(ox - sx, sy - oy))) < 3


class TestTheCombatArea:
    def test_it_is_framed_tighter_than_the_overview(self, tmp_path: Path) -> None:
        inputs = _inputs(tmp_path)
        overview = render_overview_map(**inputs, out=tmp_path / "o.png", cache_dir=tmp_path, fetch=_white, width=1000)
        combat = render_overview_map(
            **inputs, out=tmp_path / "c.png", cache_dir=tmp_path, fetch=_white, width=1000, combat_area=True
        )
        assert combat.metres_per_pixel < overview.metres_per_pixel
        assert "Khobi" in combat.symbols

    def test_its_page_follows_the_overview_and_it_is_a_dcs_picture(self, built: Built) -> None:
        report = built.report
        titles = [slide[0] for slide in _slides(report.deck)[1:]]
        tactical = titles.index("3. Situation tactique")
        assert titles[tactical + 1] == "4. Zone des combats"
        assert (built.folder / "docs" / "maps" / "combat-area.png").is_file()
        assert f"{PICTURE_PREFIX}combat-area" in report.pictures
        assert report.kneeboard[1] == "02-briefing-combat-area.jpg"


class TestTheCodeWords:
    def test_a_page_in_the_deck_and_on_the_kneeboard(self, tmp_path: Path) -> None:
        objective_mission(tmp_path)
        report = _build(tmp_path, BriefingContent(**{**CONTENT.__dict__, "code_words": CODES}))
        text = "\n".join("\n".join(slide) for slide in _slides(report.deck))
        # a table since FIX-BRIEFING-PARITY-MARTEAU 07: the word, what it means, who acts on it
        assert "PORTE OUVERTE\nLe SA-6 de Sa'sa' est détruit.\nBlade franchit le col." in text
        assert report.kneeboard[-1].endswith("-briefing-code-words.jpg")

    def test_none_no_page(self, built: Built) -> None:
        report = built.report
        assert not any("Mots-codes" in slide[0] for slide in _slides(report.deck))
        assert not any("code-words" in name for name in report.kneeboard)


def _grid_to_magnetic(theatre: str, when: date, at: tuple[float, float], heading: float) -> int:
    """A grid heading made magnetic independently: GeographicLib's azimuth to a point along it, less the WMM."""
    turn = math.radians(heading)
    start = xy_to_latlon(theatre, *at)
    end = xy_to_latlon(theatre, at[0] + 1000 * math.cos(turn), at[1] + 1000 * math.sin(turn))
    true = float(Geodesic.WGS84.Inverse(*start, *end)["azi1"])
    return round(true - declination(*start, when)) % 360 or 360


class TestTheReconPhotos:
    _SIDE = {"zone": "khobi", "title": "Dépôt de Khobi", "heading": 95, "pitch": -12, "altitude_m": 850}

    def test_the_sidecar_is_read_and_the_strays_are_warned(self, tmp_path: Path) -> None:
        folder = tmp_path / "recon"
        _photo(folder, "02-b", {**self._SIDE, "order": 2})
        _photo(folder, "01-a", {**self._SIDE, "order": 3})
        _photo(folder, "03-lonely", None)
        _photo(folder, "04-elsewhere", {"zone": "Nowhere", "title": "x"})
        _photo(folder, "05-untitled", {"zone": "Khobi"})
        (folder / "06-orphan.json").write_text("{}", encoding="utf-8")
        (folder / "recon.json").write_text('{"shots": []}', encoding="utf-8")  # the shooting tool's index
        photos, warnings = load_recon(folder, ["Khobi"])
        assert [photo.image.name for photo in photos["Khobi"]] == ["02-b.jpg", "01-a.jpg"]
        assert photos["Khobi"][0].heading == 95
        joined = "\n".join(warnings)
        for stray in ("03-lonely", "Nowhere", "05-untitled", "06-orphan"):
            assert stray in joined
        assert len(warnings) == 4

    def test_no_folder_no_photo(self, tmp_path: Path) -> None:
        assert load_recon(tmp_path / "nope", ["Khobi"]) == ({}, [])

    def test_the_caption_says_what_the_sidecar_gives(self) -> None:
        _, zx, zy, _ = KHOBI_ZONE
        photo = ReconPhoto(Path("a.jpg"), "Khobi", "Dépôt", heading=5, pitch=-12, x=zx, y=zy, altitude_m=1000)
        when = date(2016, 6, 1)
        with_all = caption("Caucasus", when, photo, (0.0, 0.0))
        assert with_all.startswith("Dépôt — ")
        assert f"{_grid_to_magnetic('Caucasus', when, (zx, zy), 5):03d}°M" in with_all
        assert "-12°" in with_all and "3 281 ft" in with_all and "N" in with_all
        assert caption("Caucasus", when, ReconPhoto(Path("a.jpg"), "Khobi", "Dépôt"), (zx, zy)) == "Dépôt"

    def test_a_heading_with_no_camera_position_is_taken_at_the_objective(self) -> None:
        _, zx, zy, _ = KHOBI_ZONE
        when = date(2016, 6, 1)
        text = caption("Caucasus", when, ReconPhoto(Path("a.jpg"), "Khobi", "Dépôt", heading=95), (zx, zy))
        assert text.endswith(f" {_grid_to_magnetic('Caucasus', when, (zx, zy), 95):03d}°M")  # "cap" or "heading"

    def test_they_follow_their_zoom_two_to_a_slide_and_a_kneeboard_page(self, tmp_path: Path, built: Built) -> None:
        objective_mission(tmp_path)
        plain = built.report
        for stem in ("a", "b", "c"):
            _photo(tmp_path / "docs" / "recon", stem, self._SIDE)
        report = _build(tmp_path)
        assert [path.name for path in report.recon] == ["a.jpg", "b.jpg", "c.jpg"]
        slides = _slides(report.deck)
        titles = [slide[0] for slide in slides]
        zoom = titles.index("5. KHOBI — Détruire le dépôt de Khobi")
        assert titles[zoom + 1 : zoom + 3] == ["Reconnaissance — KHOBI"] * 2
        assert any(re.search(r"cap \d{3}°M", text) for text in slides[zoom + 1])
        assert report.pages == plain.pages + 2
        assert report.pdf_pages == plain.pdf_pages + 2
        recon_pages = [name for name in report.kneeboard if "recon" in name]
        assert recon_pages == ["04-briefing-objective-khobi-recon-1.jpg", "05-briefing-objective-khobi-recon-2.jpg"]

    def test_a_recon_folder_elsewhere(self, tmp_path: Path) -> None:
        objective_mission(tmp_path / "m")
        _photo(tmp_path / "photos", "a", self._SIDE)
        report = _build(tmp_path / "m", recon_folder=tmp_path / "photos")
        assert [path.name for path in report.recon] == ["a.jpg"]

    def test_a_photo_page_keeps_the_photo_s_proportions(self, tmp_path: Path) -> None:
        from mission_briefing.briefing_kneeboard import photo_page

        image = _photo(tmp_path, "tall", None, size=(400, 800))
        page = Image.open(io.BytesIO(photo_page("T", [(image, "x")])))
        assert page.size == (1536, 2304)
