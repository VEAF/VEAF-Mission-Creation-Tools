"""`build_mission_briefing` — the briefing of a one-shot mission, for an assistant with no shell (FEAT-MCP-MISSION-BRIEFING 05)."""

from __future__ import annotations

from pathlib import Path

import pytest
from objective_mission_fixture import objective_mission
from veaf_mission_mcp.actions import register_default_actions
from veaf_mission_mcp.catalog import ActionCatalog

PARAMS = {
    "title": "Marteau de Khobi",
    "context": ["Les dépôts de Khobi ravitaillent le front."],
    "mission_bullets": ["Détruire le dépôt"],
    "objectives": [{"zone": "Khobi", "title": "Détruire le dépôt de Khobi", "task": ["Le dépôt et son SA-8."]}],
    "threats": [{"name": "SA-6 Gori", "group": "SA-6 Site", "family": "SA-6", "engagement_nm": 13, "source": "jeu"}],
    "pilots": {"Kobuleti Viper": ["Zip"]},
}


@pytest.fixture
def catalog() -> ActionCatalog:
    catalog = ActionCatalog()
    register_default_actions(catalog)
    return catalog


def test_the_action_runs_through_the_catalog(catalog: ActionCatalog, tmp_path: Path) -> None:
    """The wiring, not the worker: what the assistant actually calls."""
    objective_mission(tmp_path)
    result = catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS})
    assert Path(result["deck"]) == tmp_path / "docs" / "briefing.pptx" and Path(result["deck"]).is_file()
    # with the flight plan's two step sheets (FEAT-OBJECTIVE-BRIEFING-QUALITY 05) and the combat area
    assert result["pages"] == 12
    assert [Path(path).name for path in result["maps"]] == ["overview.png", "combat-area.png", "objective-khobi.png"]
    assert result["pictures"] == [
        "MCP_Briefing_overview",
        "MCP_Briefing_combat-area",
        "MCP_Briefing_objective-khobi",
    ]
    assert result["map_offline"] is True  # no network in the tests
    assert any("plain background" in warning for warning in result["warnings"])


def test_code_words_and_recon_photos_reach_the_briefing(catalog: ActionCatalog, tmp_path: Path) -> None:
    """FEAT-OBJECTIVE-MISSION-V2 03-04, through the catalog: the parameters the assistant passes."""
    import json

    from PIL import Image

    objective_mission(tmp_path / "m")
    photos = tmp_path / "photos"
    photos.mkdir()
    Image.new("RGB", (800, 450), (90, 110, 70)).save(photos / "a.jpg")
    (photos / "a.json").write_text(json.dumps({"zone": "Khobi", "title": "Le dépôt", "heading": 95}), "utf-8")
    result = catalog.run_action(
        "build_mission_briefing",
        {
            "mission_path": str(tmp_path / "m"),
            **PARAMS,
            "code_words": [{"word": "PORTE OUVERTE", "meaning": "Le SA-6 est détruit.", "who": "Blade"}],
            "recon_folder": str(photos),
        },
    )
    assert [Path(path).name for path in result["recon"]] == ["a.jpg"]
    assert any("code-words" in name for name in result["kneeboard"])
    assert any("recon" in name for name in result["kneeboard"])


def test_topo_is_a_basemap_and_anything_else_is_refused(catalog: ActionCatalog, tmp_path: Path) -> None:
    objective_mission(tmp_path)
    catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS, "basemap": "topo"})
    with pytest.raises(ValueError, match="basemap"):
        catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS, "basemap": "google"})


def test_an_unaccented_french_briefing_is_warned_and_an_accented_one_is_not(
    catalog: ActionCatalog, tmp_path: Path
) -> None:
    """FIX-DEMO-VIDEO-FINDINGS 06: the words a player reads keep their accents."""
    objective_mission(tmp_path)
    plain = {
        **PARAMS,
        "context": ["La colonne blindee a force le defile et elle est arretee faute de carburant, pres du noeud."],
    }
    result = catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **plain})
    assert any(warning.startswith("context:") and "accent" in warning for warning in result["warnings"])

    result = catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS})
    assert not any("accent" in warning for warning in result["warnings"])


def test_the_inspection_shows_the_date_and_the_briefing_pictures(catalog: ActionCatalog, tmp_path: Path) -> None:
    """FIX-DEMO-VIDEO-FINDINGS 14: the AI opened the archive with a shell to see both."""
    from veaf_mission_mcp.inspect_built import inspect_built_mission

    objective_mission(tmp_path)
    before = inspect_built_mission(tmp_path)
    assert before["briefing_pictures"] == "absent"
    assert before["weather"][0].startswith("main: 20"), "the date comes first: YYYY-MM-DD HH:MM"

    catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS})
    after = inspect_built_mission(tmp_path)
    assert any(
        line.startswith("MCP_Briefing_overview -> briefing-overview.jpg: in the archive, for blue, neutral")
        for line in after["briefing_pictures"]
    )


def test_diversions_are_spelt_as_dcs_and_an_unknown_one_is_refused(catalog: ActionCatalog, tmp_path: Path) -> None:
    """FIX-COWORK-SESSION-FINDINGS 04: the assistant names the diversion fields, checked against the map."""
    objective_mission(tmp_path)
    result = catalog.run_action(
        "build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS, "diversions": ["kutaisi"]}
    )
    assert Path(result["deck"]).is_file()
    with pytest.raises(ValueError, match="no airfield named Atlantis"):
        catalog.run_action(
            "build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS, "diversions": ["Atlantis"]}
        )


def test_the_kneeboard_and_the_archives_are_reported(catalog: ActionCatalog, tmp_path: Path) -> None:
    """FIX-COWORK-SESSION-FINDINGS 01 and 08."""
    miz = objective_mission(tmp_path)
    result = catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS})
    assert result["archives"] == [str(miz)]
    assert result["kneeboard"][0] == "01-briefing-overview.jpg"
    result = catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS, "kneeboard": False})
    assert result["kneeboard"] == []


def _deck_text(deck: str) -> str:
    from pptx import Presentation

    return "\n".join(
        shape.text_frame.text for slide in Presentation(deck).slides for shape in slide.shapes if shape.has_text_frame
    )


def test_the_flights_on_dynamic_slots_reach_the_ato(catalog: ActionCatalog, tmp_path: Path) -> None:
    """FEAT-OBJECTIVE-BRIEFING-QUALITY 01: the callsigns, never the base's every type."""
    objective_mission(tmp_path)
    flights = [{"callsign": "Ninja", "base": "Kobuleti", "aircraft": "F-16C", "count": 2, "task": "SEAD"}]
    result = catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **PARAMS, "flights": flights})
    assert "Ninja — 2 × F-16C — SEAD (slot dyn)" in _deck_text(result["deck"])


def test_a_text_given_as_one_string_is_split_into_paragraphs_and_a_mission_string_is_prose(
    catalog: ActionCatalog, tmp_path: Path
) -> None:
    """FEAT-OBJECTIVE-BRIEFING-QUALITY 04: prose, one paragraph per blank-line-separated block."""
    objective_mission(tmp_path)
    params = {
        **PARAMS,
        "strategic_situation": "Le front est figé.\n\nLa coalition frappe en profondeur.",
        "mission_bullets": "Aveugler la défense, puis frapper le PC.",
    }
    result = catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **params})
    text = _deck_text(result["deck"])
    assert "Le front est figé.\nLa coalition frappe en profondeur." in text
    assert "Aveugler la défense" in text and "– Aveugler" not in text


def test_the_deployment_and_the_waypoint_notes_reach_the_deck(catalog: ActionCatalog, tmp_path: Path) -> None:
    """FEAT-OBJECTIVE-BRIEFING-QUALITY 05; a note for a waypoint the route lacks is warned."""
    objective_mission(tmp_path)
    params = {
        **PARAMS,
        "deployment": "Viper décolle le premier.",
        "waypoint_notes": {"KHOBI": "Largage sur le dépôt.", "NOWHERE": "x"},
    }
    result = catalog.run_action("build_mission_briefing", {"mission_path": str(tmp_path), **params})
    text = _deck_text(result["deck"])
    assert "Déploiement et coordination" in text and "Viper décolle le premier." in text
    assert "Largage sur le dépôt." in text
    assert any("NOWHERE" in warning for warning in result["warnings"])
